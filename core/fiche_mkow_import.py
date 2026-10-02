# -*- coding: utf-8 -*-
"""قراءة بطاقة شخص من ملفّ (Excel / Word) لتعمير «إضافة إسم جديد».

Le document type est la «بطاقة إرشادات حول المكوّن» : une grille de
libellés («إسم و لقب المكوّن», «رقم بطاقة التعريف الوطنية»…) avec, à côté de
chacun, la valeur. On lit la GRILLE, pas une suite de lignes :

  • pour chaque case reconnue comme libellé, la valeur est la première case
    non vide qui la suit dans le sens de lecture de la feuille, jusqu'au
    libellé suivant — si bien que « رقم بطاقة التعريف | 08089763 | الصادرة
    بتاريخ | 13/01/2026 » donne bien deux valeurs sur une même ligne ;
  • le sens (droite→gauche ou gauche→droite) est DÉDUIT de la feuille : c'est
    celui qui associe une valeur au plus grand nombre de libellés.

Trois règles, celles de toute la منظومة :
  1. On ne devine pas : un libellé inconnu est ignoré.
  2. On n'enregistre rien : le navigateur pré-remplit, l'agent relit.
  3. On ne perd rien en silence : un zéro de tête ajouté, un numéro de compte
     qu'Excel a pu arrondir, une رتبة inconnue — chacun est SIGNALÉ.
"""

import io
import logging as _logging
import os
import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, datetime

from core.importation import _normaliser
from core import bataqa_import as _B

_log = _logging.getLogger('formation.' + __name__)


class ErreurFiche(Exception):
    """Message prêt à afficher, en arabe."""


# ─── Libellés reconnus → خانة du formulaire ─────────────────────────────────
#
# Convention de la منظومة : le nom affiché partout est `nom + ' ' + prenom`.
# Un « الاسم واللقب » écrit d'un seul tenant est donc coupé au PREMIER espace :
# le premier mot va dans `nom`, le reste dans `prenom` — et le nom se réaffiche
# exactement tel qu'il était écrit, quel que soit l'ordre adopté par le fichier.

LIBELLES = (
    ('nom_complet',        ('إسم و لقب المكون', 'اسم ولقب المكون', 'الاسم واللقب', 'الإسم واللقب',
                            'الاسم و اللقب', 'الإسم و اللقب', 'اللقب والاسم', 'اللقب و الاسم',
                            'الاسم الكامل', 'اسم المكون', 'الإسم و اللقب للمكون')),
    ('nom',                ('الاسم', 'الإسم', 'الاسم الأول')),
    ('prenom',             ('اللقب', 'اللقب العائلي', 'الاسم العائلي')),
    ('grade',              ('الرتبة', 'الرّتبة')),
    ('cin',                ('رقم بطاقة التعريف الوطنية', 'رقم بطاقة التعريف', 'بطاقة التعريف الوطنية',
                            'بطاقة التعريف', 'ب ت و', 'رقم ب ت و')),
    ('cin_date',           ('الصادرة بتاريخ', 'تاريخ الإصدار', 'تاريخ اصدار بطاقة التعريف',
                            'تاريخ الإصدار بطاقة التعريف')),
    ('identifiant_unique', ('رقم المعرف الوحيد', 'المعرف الوحيد', 'المعرّف الوحيد', 'رقم المعرف',
                            'المعرف', 'المعرف الفريد')),
    ('adresse',            ('العنوان الشخصي', 'العنوان')),
    ('telephone_adm',      ('الهاتف الإداري', 'الهاتف الاداري', 'هاتف العمل', 'الهاتف القار')),
    ('telephone_gsm',      ('الهاتف الجوال', 'الهاتف الجوّال', 'رقم الهاتف الجوال', 'الجوال',
                            'الهاتف', 'رقم الهاتف')),
    ('diplome',            ('الشهادة العلمية', 'الشهادة', 'المؤهل العلمي', 'آخر شهادة علمية')),
    ('degre',              ('الدرجة', 'الرتبة المهنية')),
    ('plan_fonctionnel',   ('الخطة الوظيفية', 'الخطة')),
    ('specialite',         ('التخصص', 'الاختصاص', 'الإختصاص')),
    ('administration',     ('الإدارة', 'الادارة', 'الإدارة الأصلية')),
    ('ministere',          ('وزارة الإشراف', 'وزارة الاشراف', 'الوزارة')),
    ('lieu_travail',       ('مكان العمل', 'مقر العمل', 'مركز العمل')),
    ('jiha_marjiiya',      ('الجهة المرجعية',)),
    ('email',              ('البريد الإلكتروني', 'البريد الالكتروني', 'البريد')),
    ('banque',             ('البنك', 'اسم البنك', 'المؤسسة البنكية', 'البنك أو البريد')),
    ('agence',             ('الفرع', 'الوكالة')),
    ('num_compte',         ('رقم الحساب البنكي / البريدي', 'رقم الحساب البنكي او البريدي',
                            'رقم الحساب البنكي', 'رقم الحساب البريدي', 'رقم الحساب',
                            'الهوية البنكية', 'الحساب البنكي', 'rib')),
)

#: Libellés d'autres rubriques de la fiche : ils ne portent pas de donnée de
#: la personne, mais ils ARRÊTENT la recherche d'une valeur (jamais pris pour
#: la valeur du libellé qui les précède).
_ARRETS = frozenset(_normaliser(x) for x in (
    'تكوين مستمر', 'تكوين تأهيلي', 'عدد الساعات المبرمجة', 'عدد الساعات المنجزة',
    'إمضاء المكون', 'امضاء المكون', 'خاص بوحدة التكوين', 'وحدة التكوين', 'الملاحظات',
    'ملاحظات', 'بطاقة إرشادات حول المكون', 'التوقيع', 'الإمضاء', 'التاريخ', 'حرر ب',
))

_MAX_LIBELLE = 48

_TABLE = []
for _champ, _libs in LIBELLES:
    for _l in _libs:
        _TABLE.append((_normaliser(_l), _champ))
# Le PLUS LONG d'abord : «الهاتف الإداري» ne doit pas être pris pour «الهاتف».
_TABLE.sort(key=lambda x: -len(x[0]))


def _propre(t):
    t = _B._nettoyer(str(t or '').replace('\n', ' ').replace('\r', ' '))
    return t.strip().strip('*').strip()


def _classer(texte):
    """(خانة | 'ARRET' | None, reste) — reste = valeur écrite dans la même
    case après « : »."""
    brut = _propre(texte)
    if not brut:
        return None, ''
    tete, reste = brut, ''
    for sep in (':', '：'):
        if sep in brut:
            g, d = brut.split(sep, 1)
            if _classer(g)[0]:
                tete, reste = g, d.strip()
            break
    tete = _B._RE_PUCE.sub('', tete).strip().rstrip(':：*').strip()
    tete = re.sub(r'\([^)]*\)', '', tete).strip()           # « (إجباري) »…
    if not tete or len(tete) > _MAX_LIBELLE:
        return None, ''
    # Correspondance EXACTE : « البنك الوطني الفلاحي » commence par le libellé
    # « البنك » mais c'est une VALEUR. La liste porte les variantes réelles.
    n = _normaliser(tete)
    for cle, champ in _TABLE:
        if n == cle:
            return champ, reste
    if n in _ARRETS:
        return 'ARRET', ''
    return None, ''


def est_libelle(texte):
    """Pour le découpage d'un .doc : la case COMMENCE-t-elle par un libellé
    de DONNÉE ? (Un titre comme «بطاقة إرشادات حول المكوّن» n'en est pas un :
    c'est du texte hors tableau.)"""
    ch = _classer(texte)[0]
    return bool(ch) and ch != 'ARRET'


# ─── Valeurs ─────────────────────────────────────────────────────────────────

_CHIFFRES_AR = str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹', '01234567890123456789')
_MOIS = {'جانفي': 1, 'يناير': 1, 'فيفري': 2, 'فبراير': 2, 'مارس': 3, 'أفريل': 4, 'افريل': 4,
         'أبريل': 4, 'ماي': 5, 'مايو': 5, 'جوان': 6, 'يونيو': 6, 'جويلية': 7, 'يوليو': 7,
         'أوت': 8, 'اوت': 8, 'أغسطس': 8, 'سبتمبر': 9, 'أكتوبر': 10, 'اكتوبر': 10,
         'نوفمبر': 11, 'ديسمبر': 12}


def _sans_guillemets(t):
    return str(t).strip().strip('"\'«»“”„').strip()


def lire_date(v):
    """→ ISO « AAAA-MM-JJ », ou '' si illisible. Accepte les dates Excel, le
    format tunisien JJ/MM/AAAA et « 13 جانفي 2026 »."""
    if v is None or v == '':
        return ''
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    t = _sans_guillemets(v).translate(_CHIFFRES_AR)
    m = re.search(r'(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})', t)
    if m:
        a, mo, j = map(int, m.groups())
    else:
        m = re.search(r'(\d{1,2})\s*[-/.]\s*(\d{1,2})\s*[-/.]\s*(\d{4})', t)
        if m:
            j, mo, a = map(int, m.groups())
        else:
            m = re.search(r'(\d{1,2})\s+([^\d\s]+)\s+(\d{4})', t)
            if not m or _normaliser(m.group(2)) not in {_normaliser(k) for k in _MOIS}:
                return ''
            j, a = int(m.group(1)), int(m.group(3))
            mo = next(v for k, v in _MOIS.items() if _normaliser(k) == _normaliser(m.group(2)))
    try:
        return date(a, mo, j).isoformat()
    except ValueError:
        return ''


def _chiffres(v):
    return re.sub(r'\D', '', _sans_guillemets(v).translate(_CHIFFRES_AR))


def normaliser_valeurs(brut, types=None):
    """Nettoie les valeurs lues et rend (champs, avertissements).

    `types` : pour chaque خانة, le type Python de la case d'origine (Excel
    stocke parfois un numéro comme un NOMBRE — zéros de tête perdus, 20
    chiffres arrondis)."""
    types = types or {}
    champs, avert = {}, []

    for cle, v in brut.items():
        if v is None:
            continue
        if cle == 'cin_date':
            iso = lire_date(v)
            if iso:
                champs[cle] = iso
            elif str(v).strip():
                avert.append(f'تاريخ إصدار بطاقة التعريف «{v}» غير مقروء: أدخله يدويًّا.')
            continue
        t = _sans_guillemets(v)
        if not t:
            continue
        if cle == 'cin':
            c = _chiffres(t)
            if types.get(cle) in (int, float) and 0 < len(c) < 8:
                avert.append(f'رقم بطاقة التعريف «{c}» مخزَّن كعدد في إكسال: '
                             f'أُعيد إليه الصفر الأوّل ({c.zfill(8)}). تثبّت منه.')
                c = c.zfill(8)
            if c and len(c) != 8:
                avert.append(f'رقم بطاقة التعريف «{c}» لا يتكوّن من 8 أرقام: يجب تصحيحه قبل الحفظ.')
            champs[cle] = c or t
        elif cle == 'num_compte':
            c = _chiffres(t)
            if types.get(cle) in (int, float) and len(c) >= 16:
                avert.append('رقم الحساب مخزَّن كعدد في إكسال، وقد يكون فقد أرقامه الأخيرة '
                             '(إكسال لا يحفظ أكثر من 15 رقمًا في العدد): قارنه بالوثيقة الأصليّة.')
            if c and len(c) != 20:
                avert.append(f'رقم الحساب يتكوّن من {len(c)} رقمًا والمطلوب 20 رقمًا بالضبط: يجب تصحيحه قبل الحفظ.')
            champs[cle] = c or t
        elif cle in ('telephone_gsm', 'telephone_adm'):
            champs[cle] = re.sub(r'[^\d+ ]', '', t.translate(_CHIFFRES_AR)).strip() or t
        elif cle == 'identifiant_unique':
            champs[cle] = _chiffres(t) or t
        elif cle == 'email':
            champs[cle] = t.replace(' ', '')
        else:
            champs[cle] = t

    # « الاسم واللقب » d'un seul tenant → nom (1er mot) + prenom (le reste)
    complet = champs.pop('nom_complet', '')
    if complet and not champs.get('nom'):
        mots = complet.split()
        champs['nom'] = mots[0]
        if len(mots) > 1 and not champs.get('prenom'):
            champs['prenom'] = ' '.join(mots[1:])
    # مكان العمل absent de la fiche : la الإدارة en tient lieu, à vérifier.
    if not champs.get('lieu_travail') and champs.get('administration'):
        champs['lieu_travail'] = champs['administration']
        avert.append('مكان العمل غير مذكور في الملفّ: عُمِّر من خانة «الإدارة»، تثبّت منه.')
    return champs, avert


# ─── Grilles ─────────────────────────────────────────────────────────────────

def _paires(grille, sens):
    """grille : lignes de cases (texte, valeur brute). Rend {خانة: (valeur, type)}.

    Pour chaque libellé, on avance dans le `sens` (+1/-1) : cases vides
    sautées (cases fusionnées), arrêt au premier libellé ou marqueur."""
    res = {}
    for ligne in grille:
        n = len(ligne)
        for i, (txt, brut) in enumerate(ligne):
            ch, reste = _classer(txt)
            if not ch or ch == 'ARRET' or ch in res:
                continue
            if reste:
                res[ch] = (reste, str)
                continue
            j = i + sens
            while 0 <= j < n:
                t2, b2 = ligne[j]
                if _propre(t2):
                    if _classer(t2)[0]:
                        break
                    res[ch] = (b2 if b2 not in (None, '') else t2, type(b2))
                    break
                j += sens
    return res


def _meilleur_sens(grille, prefere=+1):
    a, b = _paires(grille, +1), _paires(grille, -1)
    if len(b) > len(a) or (len(b) == len(a) and prefere == -1):
        return b
    return a


def _depuis_grilles(grilles, prefere=+1):
    res = {}
    for g in grilles:
        for k, v in _meilleur_sens(g, prefere).items():
            res.setdefault(k, v)
    return res


# ─── Formats ─────────────────────────────────────────────────────────────────

def _lire_xlsx(data):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        _log.warning('fiche xlsx illisible', exc_info=True)
        raise ErreurFiche('تعذّرت قراءة ملفّ إكسال.')
    grilles, prefere = [], +1
    for ws in wb.worksheets:
        masquees = set()
        for plage in ws.merged_cells.ranges:
            for r in range(plage.min_row, plage.max_row + 1):
                for c in range(plage.min_col, plage.max_col + 1):
                    if (r, c) != (plage.min_row, plage.min_col):
                        masquees.add((r, c))
        g = []
        for r in range(1, (ws.max_row or 0) + 1):
            ligne = []
            for c in range(1, (ws.max_column or 0) + 1):
                if (r, c) in masquees:
                    ligne.append(('', None))
                    continue
                v = ws.cell(r, c).value
                if isinstance(v, float) and v.is_integer():
                    txt = str(int(v))
                elif isinstance(v, (datetime, date)):
                    txt = v.strftime('%Y-%m-%d')
                else:
                    txt = '' if v is None else str(v)
                ligne.append((txt, v))
            g.append(ligne)
        grilles.append(g)
        try:
            if ws.sheet_view.rightToLeft:
                prefere = +1       # colonne A à droite : la valeur suit à gauche
        except Exception:
            pass
    return _depuis_grilles(grilles, prefere)


def _lire_xls(data):
    try:
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
    except ImportError:
        raise ErreurFiche('لقراءة .xls القديم ثبّت المكتبة xlrd، أو احفظ الملفّ بصيغة .xlsx.')
    except Exception:
        raise ErreurFiche('تعذّرت قراءة ملفّ .xls.')
    grilles = []
    for sh in wb.sheets():
        g = []
        for r in range(sh.nrows):
            ligne = []
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if sh.cell_type(r, c) == 3:          # date Excel
                    try:
                        v = datetime(*xlrd.xldate_as_tuple(v, wb.datemode))
                    except Exception:
                        pass
                txt = (str(int(v)) if isinstance(v, float) and v.is_integer()
                       else v.strftime('%Y-%m-%d') if isinstance(v, datetime) else str(v))
                ligne.append((txt, v))
            g.append(ligne)
        grilles.append(g)
    return _depuis_grilles(grilles)


def _lire_docx(data):
    try:
        root = ET.fromstring(zipfile.ZipFile(io.BytesIO(data)).read('word/document.xml'))
    except Exception:
        raise ErreurFiche('تعذّرت قراءة ملفّ وورد .docx.')
    grilles = []
    for g in _B._grilles_docx(root):
        grilles.append([[(c[0], c[0]) for c in ligne] for ligne in g])
    res = _depuis_grilles(grilles)
    # « libellé : valeur » en texte libre
    for p in _B._paras_docx(root):
        ch, reste = _classer(p)
        if ch and ch != 'ARRET' and reste and ch not in res:
            res[ch] = (reste, str)
    return res


def _lire_doc(data):
    try:
        texte = _B._nettoyer_word(_B._texte_word97(data))
    except _B.ErreurBataqa as e:
        raise ErreurFiche(str(e))
    except Exception:
        raise ErreurFiche('تعذّرت قراءة ملفّ .doc. احفظه بصيغة .docx ثمّ أعد المحاولة.')
    # Chaque ligne du tableau = cases jusqu'à la marque de fin de ligne (case
    # vide). Le texte hors tableau est découpé en paragraphes.
    grilles, ligne = [[]], []
    for e in _B._elements_doc(texte, est_entete=est_libelle):
        if e[0] == 'hors':
            for l in e[1]:
                if _propre(l):
                    grilles[0].append([(l, l)])
            continue
        if _B._vide(e):
            if ligne:
                grilles[0].append(ligne)
            ligne = []
            continue
        ligne.append((e[1], e[1]))
    if ligne:
        grilles[0].append(ligne)
    res = _depuis_grilles(grilles)
    for l in texte.split('\r'):
        ch, reste = _classer(l)
        if ch and ch != 'ARRET' and reste and ch not in res:
            res[ch] = (reste, str)
    return res


def analyser(data, filename):
    """octets + nom de fichier → (champs, avertissements). N'écrit rien."""
    ext = os.path.splitext(filename or '')[1].lower()
    lecteurs = {'.xlsx': _lire_xlsx, '.xlsm': _lire_xlsx, '.xls': _lire_xls,
                '.docx': _lire_docx, '.doc': _lire_doc}
    if ext not in lecteurs:
        raise ErreurFiche('صيغة غير مدعومة. حمّل ملفّ .doc أو .docx أو .xls أو .xlsx.')
    lu = lecteurs[ext](data)
    brut = {k: v for k, (v, _t) in lu.items()}
    types = {k: t for k, (_v, t) in lu.items()}
    return normaliser_valeurs(brut, types)


def rapprocher_grade(grade, grades):
    """« نقيب » → « النقيب » si c'est la graphie de la liste ; '' sinon."""
    if not grade:
        return ''
    def forme(x):
        n = _normaliser(x)
        n = re.sub(r'\s+لل?ديوان[هة]?$', '', n)
        return n[2:] if n.startswith('ال') else n
    cible = forme(grade)
    for g in grades or []:
        if forme(g) == cible:
            return g
    return ''
