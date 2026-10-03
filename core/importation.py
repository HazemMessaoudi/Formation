# -*- coding: utf-8 -*-
"""توريد المكوّنين والمشاركين من ملفّ إكسال.

Ce module lit un classeur et rend un *plan* : ce qu'il faudrait écrire, ce
qu'il faudrait compléter, ce qu'il faut refuser et pourquoi. Il n'écrit rien
lui-même et n'importe ni Flask ni la قاعدة — d'où la possibilité de
l'éprouver ligne à ligne, ce qui est exactement ce qu'on veut d'un code qui
touchera un jour à une liste de deux cents agents.

Trois règles gouvernent tout le reste :

1. **المعرّف الوحيد مفتاح.** Deux personnes ne partagent pas un معرّف ;
   deux lignes qui le partagent sont la même personne saisie deux fois.

2. **On n'écrase rien.** Une بطاقة déjà en قائمة n'est jamais réécrite par
   le fichier. Au mieux — et sur demande explicite — ses خانات restées vides
   sont complétées. Un fichier d'importation est un apport, pas une autorité.

3. **On ne devine pas.** Une colonne non reconnue est ignorée et dite comme
   telle ; un nom en une seule colonne n'est pas découpé en اسم et لقب au
   jugé. Le système n'ajoute ni ne retranche un caractère.
"""

import logging as _logging
_log = _logging.getLogger('formation.' + __name__)

import re
import unicodedata

from core import validation as _val


class ErreurImportation(Exception):
    """Le classeur n'a pas pu être lu ou ne porte pas d'entête reconnaissable."""


# ─── الأعمدة المعروفة ────────────────────────────────────────────────────────
#
# Le premier libellé de chaque couple est celui qu'écrit le نموذج engendré par
# la منظومة ; les suivants sont les graphies rencontrées ailleurs. La
# comparaison est faite après normalisation (voir `_normaliser`), si bien que
# « المعرّف الوحيد », « المعرف الوحيد » et « المُعرِّف الوحيد: » se valent.

COLONNES = (
    ('identifiant_unique', ('المعرّف الوحيد', 'المعرف', 'المعرف الفريد',
                            'الرقم التعريفي', 'م.و')),
    ('grade',              ('الرتبة', 'الرّتبة')),
    # ⚠ L'ordre est celui de la منظومة, et il n'est pas arbitraire : le nom
    # affiché partout — entêtes des tableaux, القائمة الإسميّة, البطاقة — est
    # «الاسم واللقب», construit comme `nom + ' ' + prenom`. Donc `nom` porte
    # l'اسم et `prenom` porte le لقب. Inverser ces deux lignes inverserait
    # tous les noms importés.
    ('nom',                ('الاسم', 'الإسم')),
    ('prenom',             ('اللقب', 'اللّقب')),
    ('nom_complet',        ('الاسم واللقب', 'الإسم واللقب', 'اللقب والاسم',
                            'الاسم الكامل', 'الاسم و اللقب')),
    # v1.7.1 : « مكان العمل الحالي » (fichier réel du centre) n'était pas
    # reconnu → مكان العمل restait vide pour toutes les fiches importées.
    ('lieu_travail',       ('مكان العمل', 'مقر العمل', 'الخطة', 'مركز العمل',
                            'مكان العمل الحالي', 'مقر العمل الحالي', 'مكان التعيين',
                            'مقر التعيين', 'التعيين', 'المصلحة', 'مكان العمل الأصلي')),
    ('sexe',               ('الجنس',)),
    ('fiaa_omria',         ('الفئة العمرية', 'الفئة العمريّة')),
    ('jiha_marjiiya',      ('الجهة المرجعية', 'الجهة', 'الإدارة الجهوية')),
    ('administration',     ('الإدارة',)),
    ('telephone_gsm',      ('الهاتف', 'رقم الهاتف', 'الهاتف الجوال')),
    ('email',              ('البريد الإلكتروني', 'البريد الالكتروني', 'البريد')),
    ('cin',                ('بطاقة التعريف', 'رقم بطاقة التعريف', 'ب.ت.و')),
    ('specialite',         ('الاختصاص', 'الإختصاص')),
    ('diplome',            ('الشهادة', 'الشهادة العلمية')),
)

#: Colonnes du نموذج proposé au téléchargement, dans l'ordre d'affichage.
COLONNES_MODELE = (
    'المعرّف الوحيد', 'الرتبة', 'الاسم', 'اللقب',
    'مكان العمل', 'الجهة المرجعية', 'الهاتف',
)

#: Au-delà, ce n'est plus une liste d'agents : on refuse plutôt que de ramer.
MAX_LIGNES = 5000

#: Nombre de lignes du haut examinées à la recherche de l'entête. Les tableaux
#: administratifs portent volontiers un titre, un logo et deux lignes vides
#: avant la première خانة utile.
FENETRE_ENTETE = 12


# ─── التّطبيع ────────────────────────────────────────────────────────────────

_DIACRITIQUES = re.compile(r'[ً-ْٰـ]')
_PONCTUATION = re.compile(r'[:\.\*\-_/\\()\[\]«»"\'،؛؟]')


def _normaliser(texte):
    """Rend une forme comparable : sans حركات, sans هَمَزات décoratives.

    Sert aux entêtes comme aux noms. On ne stocke jamais cette forme — elle
    ne sert qu'à reconnaître ; ce qui est écrit dans la قاعدة reste le texte
    du fichier, caractère pour caractère.
    """
    if texte is None:
        return ''
    t = unicodedata.normalize('NFKC', str(texte))
    t = _DIACRITIQUES.sub('', t)
    t = _PONCTUATION.sub(' ', t)
    t = (t.replace('أ', 'ا').replace('إ', 'ا').replace('آ', 'ا')
          .replace('ى', 'ي').replace('ة', 'ه').replace('ؤ', 'و')
          .replace('ئ', 'ي'))
    return re.sub(r'\s+', ' ', t).strip().lower()


def _texte(valeur):
    """Rend la خانة telle qu'un humain l'a écrite.

    إكسال rend volontiers `77001` en nombre flottant : sans ce passage, un
    معرّف وحيد deviendrait « 77001.0 » et ne correspondrait plus à rien.
    """
    if valeur is None:
        return ''
    if isinstance(valeur, bool):
        return ''
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    if isinstance(valeur, int):
        return str(valeur)
    return re.sub(r'\s+', ' ', str(valeur)).strip()


_LIBELLES = {}
for _champ, _libelles in COLONNES:
    for _lib in _libelles:
        _LIBELLES.setdefault(_normaliser(_lib), _champ)


#: v1.7.1 — entêtes reconnus par leur DÉBUT quand la forme exacte manque :
#: « مكان العمل الحالي », « مكان العمل (المصلحة) », « مقر العمل بالإدارة »…
_PREFIXES = tuple((_normaliser(debut), champ) for debut, champ in (
    ('مكان العمل', 'lieu_travail'), ('مقر العمل', 'lieu_travail'),
    ('مركز العمل', 'lieu_travail'), ('مكان التعيين', 'lieu_travail'),
    ('المعرف الوحيد', 'identifiant_unique'), ('الفئة العمرية', 'fiaa_omria'),
))


def champ_de_libelle(libelle):
    """Rend le nom de خانة correspondant à un entête, ou None."""
    n = _normaliser(libelle)
    champ = _LIBELLES.get(n)
    if champ or not n:
        return champ
    for debut, champ in _PREFIXES:
        if n.startswith(debut + ' '):
            return champ
    return None


def _valeur_choix(champ, valeur):
    """v1.7.1 — الجنس / الفئة العمريّة lus dans un fichier : ramenés à l'une
    des valeurs de la منظومة, sinon ignorés (jamais de valeur inventée)."""
    n = _normaliser(valeur)
    if not n:
        return ''
    if champ == 'sexe':
        if n in ('ذكر', 'm', 'h', 'homme', 'male', 'رجل'):
            return 'ذكر'
        if n in ('انثي', 'f', 'femme', 'female', 'امراه'):
            return 'أنثى'
        return ''
    if champ == 'fiaa_omria':
        from core.validation import FIAAT
        for option in FIAAT:
            if n == _normaliser(option):
                return option
        if 'اقل' in n:
            return FIAAT[0]
        if 'اكثر' in n or 'فما فوق' in n:
            return FIAAT[1]
    return ''


# ─── قراءة الملفّ ────────────────────────────────────────────────────────────

def _charger_feuille(source, feuille=None):
    """Charge la première feuille utile d'un classeur .xlsx ou .xls."""
    # Lire les octets une seule fois pour détecter le format
    if hasattr(source, 'read'):
        octets = source.read()
        source.seek(0) if hasattr(source, 'seek') else None
    else:
        octets = source
        source = None

    # Signature magic: .xls (BIFF) commence par D0 CF 11 E0
    est_xls = octets[:4] == b'\xd0\xcf\x11\xe0'

    if est_xls:
        # ── Format .xls (BIFF) — lecture via xlrd ────────────────────────
        try:
            import xlrd
        except ImportError:                                # pragma: no cover
            raise ErreurImportation('تعذّر تحميل قارئ ملفّات .xls (xlrd).')
        try:
            classeur = xlrd.open_workbook(file_contents=octets)
        except Exception:
            _log.warning('_charger_feuille : exception ignorée', exc_info=True)
            raise ErreurImportation(
                'تعذّرت قراءة الملفّ بشكل .xls — '
                'تأكّد من أنّ الملفّ ليس تالفًا.')
        noms = classeur.sheet_names()
        if feuille and feuille in noms:
            ws = classeur.sheet_by_name(feuille)
        else:
            ws = classeur.sheet_by_index(0)
        lignes = []
        for i in range(min(ws.nrows, MAX_LIGNES + FENETRE_ENTETE)):
            lignes.append([
                c.value if c.ctype not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)
                else None
                for c in ws.row(i)
            ])
        return ws.name, noms, lignes

    else:
        # ── Format .xlsx — lecture via openpyxl ──────────────────────────
        from io import BytesIO
        try:
            from openpyxl import load_workbook
        except ImportError:                                # pragma: no cover
            raise ErreurImportation('تعذّر تحميل قارئ ملفّات إكسال (openpyxl).')
        try:
            classeur = load_workbook(BytesIO(octets), read_only=True, data_only=True)
        except Exception:
            _log.warning('_charger_feuille : exception ignorée', exc_info=True)
            raise ErreurImportation(
                'تعذّرت قراءة الملفّ. يُقبل شكل .xlsx أو .xls — '
                'تأكّد من أنّ الملفّ ليس تالفًا.')
        try:
            if feuille and feuille in classeur.sheetnames:
                ws = classeur[feuille]
            else:
                ws = classeur[classeur.sheetnames[0]]
            lignes = []
            for i, ligne in enumerate(ws.iter_rows(values_only=True), start=1):
                if i > MAX_LIGNES + FENETRE_ENTETE:
                    break
                lignes.append(list(ligne))
            return ws.title, classeur.sheetnames, lignes
        finally:
            try:
                classeur.close()
            except Exception:
                pass


def _reperer_entete(lignes):
    """Trouve la ligne d'entête : celle qui reconnaît le plus de خانات.

    Un tableau administratif commence rarement par son entête. Plutôt que
    d'exiger de l'utilisateur qu'il nettoie son fichier, on cherche.
    """
    meilleur, meilleur_i, meilleur_map = 0, None, {}
    for i, ligne in enumerate(lignes[:FENETRE_ENTETE]):
        correspondances = {}
        for j, cellule in enumerate(ligne):
            champ = champ_de_libelle(_texte(cellule))
            if champ and champ not in correspondances.values():
                correspondances[j] = champ
        if len(correspondances) > meilleur:
            meilleur, meilleur_i, meilleur_map = len(correspondances), i, correspondances
    if meilleur < 2:
        raise ErreurImportation(
            'لم يُعثر على سطر عناوين في الملفّ. يجب أن يحتوي أحد أسطره الأولى '
            'على عنوانين على الأقلّ من العناوين المعروفة (مثل «المعرّف الوحيد» و«اللقب»). '
            'نزّل النّموذج المقترح واعتمد عناوينه.')
    return meilleur_i, meilleur_map


# ─── بناء الخطّة ─────────────────────────────────────────────────────────────

def _cle_nom(fiche):
    """Clé de repli quand le معرّف manque : le nom complet normalisé."""
    plein = ' '.join(x for x in (fiche.get('nom', ''), fiche.get('prenom', '')) if x)
    return _normaliser(plein)


def analyser(source, existantes, *, accepter_sans_identifiant=False,
             completer_existantes=False, feuille=None):
    """Lit le classeur et rend le plan de ce qui serait écrit.

    `existantes` : les بطاقات déjà en قائمة (dicts portant au moins `id`,
    `nom`, `prenom`, `identifiant_unique`).

    Rien n'est écrit ici. Le plan rendu porte :
      * `a_ajouter`    — [(n° de ligne, بيانات)]
      * `a_completer`  — [(n° de ligne, id, بيانات partielles)]
      * `rejetes`      — [{ligne, nom, identifiant, motif}]
    """
    titre_feuille, feuilles, lignes = _charger_feuille(source, feuille)
    i_entete, colonnes = _reperer_entete(lignes)

    libelles_ignores = []
    for j, cellule in enumerate(lignes[i_entete]):
        libelle = _texte(cellule)
        if libelle and j not in colonnes:
            libelles_ignores.append(libelle)

    # Index des بطاقات déjà en قائمة
    par_identifiant, par_nom = {}, {}
    for f in existantes or []:
        ident = _texte(f.get('identifiant_unique'))
        if ident:
            par_identifiant.setdefault(ident, f)
        cle = _cle_nom(f)
        if cle:
            par_nom.setdefault(cle, f)

    a_ajouter, a_completer, rejetes = [], [], []
    vus_identifiants, vus_noms = {}, {}
    nb_lignes = 0

    for decalage, ligne in enumerate(lignes[i_entete + 1:]):
        numero = i_entete + 2 + decalage          # numéro tel que l'affiche إكسال
        if nb_lignes >= MAX_LIGNES:
            break

        brut = {}
        for j, champ in colonnes.items():
            brut[champ] = _texte(ligne[j]) if j < len(ligne) else ''
        if not any(brut.values()):
            continue                              # سطر فارغ : ni accepté ni refusé
        nb_lignes += 1

        # Un nom en une seule colonne reste en une seule خانة : on ne découpe
        # pas « بن عبد الله محمد الهادي » au jugé.
        # رقم بطاقة التعريف (v1.6) : إكسال mange le zéro de tête d'un CIN
        # saisi comme nombre (01234567 → 1234567) ; on le restitue, puis la
        # forme est normalisée (8 chiffres exactement, contrôlé plus bas).
        if brut.get('cin'):
            j_cin = next((j for j, c in colonnes.items() if c == 'cin'), None)
            cellule = ligne[j_cin] if j_cin is not None and j_cin < len(ligne) else None
            v = _val.normaliser(brut['cin'])
            if isinstance(cellule, (int, float)) and not isinstance(cellule, bool) \
                    and v.isdigit() and len(v) < 8:
                v = v.zfill(8)
            brut['cin'] = v

        for _c in ('sexe', 'fiaa_omria'):
            if _c in brut:
                brut[_c] = _valeur_choix(_c, brut[_c])

        complet = brut.pop('nom_complet', '')
        if complet and not brut.get('nom'):
            brut['nom'] = complet
        fiche = {k: v for k, v in brut.items() if v}

        nom_affiche = ' '.join(x for x in (fiche.get('nom', ''),
                                           fiche.get('prenom', '')) if x)
        ident = fiche.get('identifiant_unique', '')

        def _refus(motif):
            rejetes.append({'ligne': numero, 'nom': nom_affiche,
                            'identifiant': ident, 'motif': motif})

        if not nom_affiche:
            _refus('بدون اسم')
            continue

        if fiche.get('cin') and _val.erreur_champ('cin', fiche['cin']):
            _refus('رقم بطاقة التعريف يجب أن يتكوّن من 8 أرقام')
            continue

        if not ident:
            if not accepter_sans_identifiant:
                _refus('بدون معرّف وحيد')
                continue
            cle = _cle_nom(fiche)
            if cle in vus_noms:
                _refus(f'مكرّر داخل الملفّ (السّطر {vus_noms[cle]})')
                continue
            vus_noms[cle] = numero
            deja = par_nom.get(cle)
        else:
            if ident in vus_identifiants:
                _refus(f'معرّف مكرّر داخل الملفّ (السّطر {vus_identifiants[ident]})')
                continue
            vus_identifiants[ident] = numero
            deja = par_identifiant.get(ident)

        if deja is None:
            a_ajouter.append((numero, fiche))
            continue

        if not completer_existantes:
            _refus('مسجّل مسبقا في القائمة')
            continue

        # Compléter, jamais écraser : seules les خانات vides côté قائمة sont
        # remplies. Une valeur déjà saisie par un agent prime sur le fichier.
        partiel = {k: v for k, v in fiche.items()
                   if not _texte(deja.get(k))}
        if partiel:
            a_completer.append((numero, deja.get('id'), partiel))
        else:
            _refus('مسجّل مسبقا ولا خانة فارغة تُستكمل')

    return {
        'feuille': titre_feuille,
        'feuilles': feuilles,
        'ligne_entete': i_entete + 1,
        'colonnes_reconnues': sorted(set(colonnes.values())),
        'colonnes_ignorees': libelles_ignores,
        'total_lignes': nb_lignes,
        'a_ajouter': a_ajouter,
        'a_completer': a_completer,
        'rejetes': rejetes,
    }


# ─── النّموذج المقترح ────────────────────────────────────────────────────────

def construire_modele(exemples=None):
    """Engendre le نموذج (.xlsx) à télécharger, en mémoire.

    Un modèle vaut mieux qu'une page d'explications : l'utilisateur ouvre,
    remplit, renvoie. Les عناوين y sont exactement celles que la منظومة
    reconnaît.
    """
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = 'المكوّنون'
    ws.sheet_view.rightToLeft = True

    ws.append(list(COLONNES_MODELE))
    fond = PatternFill('solid', fgColor='1E3A5F')
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color='FFFFFF', size=12)
        cellule.fill = fond
        cellule.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26
    for lettre in 'ABCDEFG':
        ws.column_dimensions[lettre].width = 22
    ws.freeze_panes = 'A2'

    for ligne in (exemples or []):
        ws.append(list(ligne))

    flux = BytesIO()
    wb.save(flux)
    flux.seek(0)
    return flux
