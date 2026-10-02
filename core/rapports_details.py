# -*- coding: utf-8 -*-
"""V3 — Phase 6b : التّقارير التّفصيليّة et export Excel.

Une période → des tableaux prêts à afficher (page « التّقارير التّفصيليّة »)
et à exporter (un classeur .xlsx, une feuille par tableau) :
المكوّنون، الدّورات، الأشهر، الجنس، الفئة العمريّة، المستوى، الصّيغة، التّعاون.
Le même dictionnaire sert aux deux : ce que l'on voit est ce que l'on exporte."""
from io import BytesIO

from core import rapports_moteur as rm
from core import exports


def _pct(a, b):
    return round(100.0 * a / b, 2) if b else 0.0


def _tab(tid, titre, entetes, lignes, total=None, note=''):
    return {'id': tid, 'titre': titre, 'entetes': list(entetes),
            'lignes': [list(l) for l in lignes], 'total': list(total) if total else None,
            'note': note}


def _fr(iso):
    return f'{iso[8:10]}/{iso[5:7]}/{iso[:4]}' if iso and len(iso) >= 10 else (iso or '')


def tableaux(periode, critere='finalisees'):
    r = rm.calculer(periode, critere)
    act, part, pres, form = r['activites'], r['participants'], r['presence'], r['formateurs']
    tabs = []

    # ── ملخّص ──
    g = pres['global']
    tabs.append(_tab('resume', 'ملخّص الفترة', ('البيان', 'القيمة'), [
        ('الأنشطة المنجزة', act['total']), ('أيّام التّكوين', act['jours']),
        *[(f"• {c['categorie']}", c['activites']) for c in act['categories']],
        ('المشاركون', part['total']), ('المشاركون دون تكرار', part['uniques']),
        ('معدّل المشاركين في النّشاط', part['moyenne']),
        ('الحاضرون', g['presents']), ('الغيابات', g['absents']),
        ('نسبة الحضور (%)', g['taux_presence']),
        ('المكوّنون', form['total']),
        ('السّاعات (مستحقّات منجزة)', act['heures']), ('المستحقّات (د)', act['montant']),
        ('نسبة إنجاز المخطّط (%)', r['realisation']['taux'] if r['realisation']['taux'] is not None else '—'),
        ('دورات غير مختومة (غير محتسبة)', r['non_finalisees'] if critere == 'finalisees' else 0),
    ]))

    # ── المكوّنون ──
    fl = r['par_formateur']
    tabs.append(_tab('formateurs', 'حسب المكوّنين',
                     ('المكوّن', 'الرّتبة', 'الصّنف', 'الجنس', 'الأنشطة', 'الأيّام', 'المشاركون',
                      'السّاعات', 'المستحقّات (د)'),
                     [(f['nom'], f['grade'], f['categorie'], f['sexe'], f['activites'], f['jours'],
                       f['participants'], f['heures'], f['montant']) for f in fl],
                     total=('المجموع', '', '', '', sum(f['activites'] for f in fl),
                            sum(f['jours'] for f in fl), sum(f['participants'] for f in fl),
                            round(sum(f['heures'] for f in fl), 2),
                            round(sum(f['montant'] for f in fl), 3)),
                     note='دورة يؤمّنها أكثر من مكوّن تُحتسب لكلّ واحد منهم؛ السّاعات والمبالغ لا تُوزَّع إلّا لمكوّن وحيد.'))

    # ── الدّورات ──
    dl = r['par_dorra']
    tabs.append(_tab('dorrat', 'حسب الدّورات',
                     ('الدّورة', 'من', 'إلى', 'الأيّام', 'المكوّن', 'التّصنيف', 'المستوى', 'الصّيغة',
                      'المشاركون', 'الحاضرون', 'الغيابات', 'نسبة الحضور (%)'),
                     [(d['titre'], _fr(d['date_formation']), _fr(d['date_fin'] or d['date_formation']),
                       d['jours'], f"{d['grade']} {d['formateur']}".strip(), d['categorie'],
                       d['niveau'], d['mode'], d['participants'], d['presents'], d['absents'],
                       d['taux_presence']) for d in dl],
                     total=('المجموع', '', '', sum(d['jours'] for d in dl), '', '', '', '',
                            sum(d['participants'] for d in dl), sum(d['presents'] for d in dl),
                            sum(d['absents'] for d in dl), g['taux_presence'])))

    # ── الأشهر ──
    pm = r['par_mois']
    tabs.append(_tab('mois', 'حسب الأشهر',
                     ('الشّهر', 'السّنة', 'الأنشطة', 'الأيّام', 'المشاركون', 'الغيابات',
                      'نسبة الحضور (%)'),
                     [(m['mois'], m['annee'], m['activites'], m['jours'], m['participants'],
                       m['absents'], _pct(m['participants'] - m['absents'], m['participants']))
                      for m in pm],
                     total=('المجموع', '', act['total'], act['jours'], part['total'], g['absents'],
                            g['taux_presence'])))

    # ── الجنس ──
    lignes = []
    for s in part['par_sexe']:
        b = pres['par_sexe'][s]
        lignes.append((s, part['par_sexe'][s], part['part_sexe'][s], part['moyenne_sexe'][s],
                       b['absents'], b['taux_presence'] if b['total'] else '—',
                       form['par_sexe'].get(s, 0)))
    tabs.append(_tab('sexe', 'حسب الجنس',
                     ('الجنس', 'المشاركون', 'النّسبة (%)', 'المعدّل في النّشاط', 'الغيابات',
                      'نسبة الحضور (%)', 'المكوّنون'), lignes,
                     total=('المجموع', part['total'], 100 if part['total'] else 0, part['moyenne'],
                            g['absents'], g['taux_presence'], form['total'])))

    # ── الفئة العمريّة ──
    tabs.append(_tab('fiaa', 'حسب الفئة العمريّة', ('الفئة العمريّة', 'المشاركون', 'النّسبة (%)'),
                     [(k, v, _pct(v, part['total'])) for k, v in part['par_fiaa'].items()],
                     total=('المجموع', part['total'], 100 if part['total'] else 0)))

    # ── الصّفة (رتب المشاركين) ──
    lignes = []
    for k in rm.CATEGORIES_PARTICIPANTS:
        b = pres['par_categorie'][k]
        lignes.append((k, part['par_categorie'][k], part['part_categorie'][k],
                       part['moyenne_categorie'][k], b['absents'],
                       b['taux_presence'] if b['total'] else '—'))
    tabs.append(_tab('categorie', 'حسب صفة المشاركين',
                     ('الصّفة', 'المشاركون', 'النّسبة (%)', 'المعدّل في النّشاط', 'الغيابات',
                      'نسبة الحضور (%)'), lignes,
                     total=('المجموع', part['total'], 100 if part['total'] else 0, part['moyenne'],
                            g['absents'], g['taux_presence'])))

    # ── المستوى / الصّيغة / التّعاون ──
    for tid, titre, cle, lib in (('niveau', 'حسب المستوى', 'par_niveau', 'المستوى'),
                                 ('mode', 'حسب الصّيغة', 'par_mode', 'الصّيغة'),
                                 ('cooperation', 'حسب التّعاون', 'par_cooperation', 'التّعاون')):
        x = r[cle]
        tabs.append(_tab(tid, titre, (lib, 'الأنشطة', 'النّسبة (%)', 'الأيّام', 'المشاركون'),
                         [(v['valeur'], v['activites'], _pct(v['activites'], act['total']), v['jours'],
                           v['participants']) for v in x],
                         total=('المجموع', sum(v['activites'] for v in x), '',
                                sum(v['jours'] for v in x), sum(v['participants'] for v in x))))
    return {'periode': periode, 'critere': critere, 'tableaux': tabs,
            'non_finalisees': r['non_finalisees']}


def classeur(d, centre=''):
    """Le classeur .xlsx : une feuille par tableau (BytesIO)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    p = d['periode']
    for t in d['tableaux']:
        ws = wb.create_sheet(t['titre'].replace('حسب ', '')[:31])
        lignes = t['lignes'] + ([t['total']] if t['total'] else [])
        sous = f"{t['titre']} — {p['libelle']}" + (f' — {centre}' if centre else '')
        exports._remplir(ws, t['entetes'], lignes, sous_titre=sous)
        if t['total']:
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
    flux = BytesIO()
    wb.save(flux)
    flux.seek(0)
    return flux
