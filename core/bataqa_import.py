# -*- coding: utf-8 -*-
"""قراءة بطاقة بيداغوجية من ملفّ (Excel/Word/قالب) لملء نموذج مادّة تكوين.

La بطاقة بيداغوجية réelle est un **tableau** (le même que la منظومة imprime) :
des cases d'en-tête (موضوع التكوين، محاور الدورة…) et, en dessous ou à côté,
la valeur. On lit donc la GRILLE, pas une suite de lignes : pour chaque case
reconnue comme un libellé, la valeur est la case juste en dessous (même
colonne) ou la case voisine sur la même ligne. Les listes à puces d'une même
case sont recollées avec des retours à la ligne.

On ne devine jamais : un libellé inconnu est ignoré. Et on ne sauvegarde
rien — l'utilisateur relit le formulaire avant d'enregistrer.

Formats acceptés : .xlsx / .xlsm (openpyxl), .xls (xlrd si présent), .docx
(zip + XML de la bibliothèque standard), .doc (extraction texte au mieux —
le .docx reste le plus fiable).
"""

import logging as _logging
_log = _logging.getLogger('formation.' + __name__)

import io
import os
import re
import zipfile
import xml.etree.ElementTree as ET

from core.importation import _normaliser


# Libellés reconnus → خانة du formulaire مادّة التكوين.
CHAMPS = (
    ('titre',                  ('موضوع التكوين', 'عنوان التكوين', 'عنوان الدورة',
                                'عنوان المادة', 'الموضوع')),
    ('type_formation',         ('نوع التكوين', 'النوع')),
    ('mustahdafun',            ('المستهدفون بالتكوين', 'المستهدفون',
                                'الفئة المستهدفة', 'الفئة المستهدفة بالتكوين')),
    ('lieu_formation_defaut',  ('مكان التكوين', 'المكان', 'مكان الدورة')),
    ('mahawer',                ('محاور الدورة', 'المحاور', 'محاور التكوين')),
    ('objectifs',              ('أهداف الدورة', 'الأهداف', 'أهداف التكوين')),
    ('methodes_pedagogiques',  ('الطرق البيداغوجية', 'الطرائق البيداغوجية',
                                'الطريقة البيداغوجية')),
    ('moyens_pedagogiques',    ('المعينات البيداغوجية', 'الوسائل والمعينات',
                                'المعينات البيداغوجيه', 'المعينات')),
    ('preparation_materielle', ('الإعداد المادي', 'التحضير المادي',
                                'الإعداد المادّي')),
    ('equipements',            ('التجهيزات والمعدات الخصوصية',
                                'التجهيزات والمعدّات الخصوصيّة',
                                'التجهيزات والمعدات', 'التجهيزات')),
)

_LIB = {}
for _c, _libs in CHAMPS:
    for _l in _libs:
        _LIB.setdefault(_normaliser(_l), _c)

TYPES_VALIDES = ('أساسي', 'مستمر', 'مستمر حضوري')

# En-têtes courants de la بطاقة بيداغوجية non mappés au formulaire.
# Ils doivent être traités comme des libellés (pas comme des valeurs) lors
# de la recherche label→valeur, sinon ils « polluent » le champ précédent.
_LIBELLES_IGNORABLES = frozenset(_normaliser(l) for l in (
    # Participants / effectif
    'عدد المشاركين', 'عدد المتكوّنين', 'عدد المتكونين',
    'عدد المشاركين والمشاركات', 'عدد المتربصين', 'عدد المتكوّنين والمتكوّنات',
    # المصالح / الجهات المشاركة  ← colonne fréquente dans la 1re grille
    'المصالح المعنية بالمشاركة', 'المصالح المشاركة', 'المصالح المعنية',
    'الجهات المشاركة', 'الجهة المشاركة',
    # Dates
    'تاريخ الدورة التكوينية', 'تاريخ الدورة', 'تاريخ التكوين',
    'تاريخ الانعقاد', 'تاريخ الإنجاز', 'تاريخ التنفيذ',
    # Durée
    'المدة الزمنية', 'مدة التكوين', 'المدة', 'المدة الزمنية للتكوين',
    # Formateur / formateurs
    'اسم المكوّن', 'اسم المكون', 'المكوّن', 'المكون', 'المكوّن/ة',
    'المكوّنون', 'المكونون', 'أسماء المكوّنين', 'أسماء المكونين',
    # Numéros / repères
    'الرقم الترتيبي', 'الرقم',
    # Observations / cadre
    'ملاحظات', 'ملاحظة',
    'الإطار العام للدرس',
    # En-têtes de section souvent captés comme cellules
    'الإطار المكاني والزماني للدورة التكوينية',
    'الإطار المكاني والزماني',
    'الطرق والمعينات البيداغوجية',
    'الإعداد المادي والتجهيزات والمعدات الخصوصية',
    'الإعداد المادي والتجهيزات',
    'بطاقة بيداغوجية',
))


_MOTIFS_LIBELLE = (
    re.compile(r'عدد\s*الم[شت]'),
    re.compile(r'تاريخ\s+ال[دط]ورة'),
    re.compile(r'تاريخ\s+التكوين'),
    re.compile(r'تاريخ\s+الانعقاد'),
    re.compile(r'تاريخ\s+(الإ|ال)نجاز'),
    re.compile(r'مدة\s+التكوين'),
    re.compile(r'المدة\s*الزمنية'),
    re.compile(r'اسم\s+المكوّ?ن'),
    re.compile(r'^المكوّ?ن'),
    re.compile(r'الإطار\s+العام\s+لل'),
    re.compile(r'الإطار\s+المكاني'),
    re.compile(r'الرقم\s+الترتيبي'),
    re.compile(r'^ملاحظا?ت?$'),
    re.compile(r'المصالح\s+المعني'),      # المصالح المعنية بالمشاركة
    re.compile(r'الجهات?\s+المشارك'),     # الجهة/الجهات المشاركة
    re.compile(r'بطاقة\s+بيداغوجية'),
    re.compile(r'الطرق\s+والمعينات\s+البيداغوجية'),
    re.compile(r'الإعداد\s+المادي\s+والتجهيزات'),
)

_RE_INVISIBLES = re.compile(r'[​-‏‪-‮﻿\xa0]+')

#: Puces et numérotations de début de ligne : « - », « • », « 1- », « 2) »,
#: « أ- », « ب) »… Elles ne font pas partie de la valeur : la بطاقة imprimée
#: pose ses propres puces.
_RE_PUCE = re.compile(r'^\s*(?:[-–—•●○▪■◆◦*·❖➢►▶✓✔◄]+|\(?\d{1,2}\s*[-.)–](?!\d)|\(?[أ-ي]\s*[-)–])\s*')

#: Un en-tête est COURT. Au-delà, c'est du contenu — même s'il contient les
#: mots d'un en-tête (« تقديم الإطار العام للدّرس » est un محور, pas le
#: libellé « الإطار العام للدرس »).
_LONGUEUR_MAX_ENTETE = 48
#: Ce qui peut suivre un libellé reconnu sans en faire du contenu :
#: « محاور الدورة التكوينية », « مكان التكوين (افتراضي) »…
_SUFFIXE_MAX = 16

IGNORE = '__ignore__'

#: Titres de SECTION de la بطاقة : ils se placent ENTRE deux tableaux. Ils ne
#: sont pas une colonne — ils annoncent qu'un nouveau tableau commence.
_SECTIONS = frozenset(_normaliser(l) for l in (
    'تقديم الدورة التكوينية', 'تقديم الدورة',
    'الإطار المكاني والزماني للدورة التكوينية', 'الإطار المكاني والزماني',
    'الطرق والمعينات البيداغوجية',
    'الإعداد المادي والتجهيزات والمعدات الخصوصية', 'الإعداد المادي والتجهيزات',
    'بطاقة بيداغوجية',
))


def _est_section(ligne):
    propre = _RE_PUCE.sub('', _nettoyer(ligne)).strip().rstrip(':：').strip()
    return bool(propre) and _normaliser(propre) in _SECTIONS


def _nettoyer(txt):
    """Supprime les caractères Unicode invisibles et normalise les espaces
    d'UNE ligne."""
    if not txt:
        return ''
    propre = _RE_INVISIBLES.sub(' ', str(txt))
    return re.sub(r'[ \t]+', ' ', propre).strip()


def _lignes(txt):
    """Les lignes non vides d'une case, nettoyées, sans puces."""
    if txt is None:
        return []
    res = []
    for ln in str(txt).replace('\r', '\n').split('\n'):
        ln = _RE_PUCE.sub('', _nettoyer(ln)).strip()
        if ln:
            res.append(ln)
    return res


class ErreurBataqa(Exception):
    """Message prêt à afficher à l'utilisateur, en arabe."""


def _correspond(n, cle):
    """`n` est-il le libellé `cle` (éventuellement suivi d'un court suffixe) ?"""
    if not n or not cle:
        return False
    if n == cle:
        return True
    return n.startswith(cle + ' ') and len(n) - len(cle) <= _SUFFIXE_MAX


def _classer_ligne(ligne):
    """Rend la خانة (ou IGNORE) si la ligne est un EN-TÊTE, sinon None.

    Jamais de devinette : une ligne longue est du contenu, point."""
    propre = _RE_PUCE.sub('', _nettoyer(ligne)).strip().rstrip(':：').strip()
    if not propre or len(propre) > _LONGUEUR_MAX_ENTETE:
        return None
    n = _normaliser(propre)
    for cle, champ in _LIB.items():
        if _correspond(n, cle):
            return champ
    if n in _LIBELLES_IGNORABLES or any(_correspond(n, i) for i in _LIBELLES_IGNORABLES):
        return IGNORE
    if n in _SECTIONS or any(_correspond(n, s) for s in _SECTIONS):
        return IGNORE
    if any(m.match(propre) for m in _MOTIFS_LIBELLE):
        return IGNORE
    return None


def _entete(texte):
    """Analyse une case (ou un paragraphe).

    Rend (champ, reste) :
      champ — la خانة si la case COMMENCE par un en-tête, IGNORE pour un
              en-tête non repris dans le formulaire, None pour du contenu ;
      reste — les lignes de contenu que la case porte après son en-tête
              (« محاور الدورة: … » ou en-tête suivi de puces dans la même case).
    """
    lignes = [l for l in str(texte or '').replace('\r', '\n').split('\n') if _nettoyer(l)]
    if not lignes:
        return None, []
    premiere = _nettoyer(lignes[0])
    # « libellé : valeur » sur la même ligne
    for sep in (':', '：'):
        if sep in premiere:
            g, d = premiere.split(sep, 1)
            ch = _classer_ligne(g)
            if ch:
                return ch, _lignes(d) + _lignes('\n'.join(lignes[1:]))
    ch = _classer_ligne(premiere)
    if ch:
        return ch, _lignes('\n'.join(lignes[1:]))
    return None, []


def _champ(libelle):
    """Compatibilité : la خانة d'un libellé, ou None."""
    ch, _ = _entete(libelle)
    return ch if ch and ch != IGNORE else None


def _est_libelle(txt):
    """True si la case commence par un en-tête (repris ou ignoré)."""
    return _entete(txt)[0] is not None


def _propre(v):
    """Valeur d'une case : lignes sans puces, recollées par retour à la ligne."""
    return '\n'.join(_lignes(v))


def _ajouter(res, champ, lignes):
    """Accumule des lignes sous une خانة, sans doublon, dans l'ordre du fichier."""
    if not champ or champ == IGNORE or not lignes:
        return
    deja = res.setdefault(champ, [])
    for l in lignes:
        if l not in deja:
            deja.append(l)


def _fusion(dest, source):
    """Deux lectures du même fichier : la lecture en GRILLE (qui sait quelle
    case est sous quel en-tête) fait foi ; une lecture linéaire ne fait que
    COMPLÉTER les خانات restées vides.

    (v1.4d préférait la valeur « la plus longue » : c'est ce qui laissait une
    lecture linéaire, qui mélange deux colonnes voisines, écraser la bonne
    valeur — «أهداف» recevait محاور + أهداف.)"""
    for k, v in source.items():
        v = (v or '').strip()
        if v and not (dest.get(k) or '').strip():
            dest[k] = v


def _finaliser(acc):
    return {k: '\n'.join(v).strip() for k, v in acc.items() if v}


# ─── Grilles (tableaux) ──────────────────────────────────────────────────────
#
# Une case = (texte, première colonne, colonne après la dernière). Les
# colonnes sont celles de la GRILLE : une case fusionnée sur deux colonnes
# couvre [c, c+2). C'est ce qui permet d'aligner un en-tête sur les cases qui
# sont vraiment sous lui, même quand Word fusionne ou découpe les cases.

def _cases(row):
    """Accepte une ligne de chaînes (ancien format) ou de triplets."""
    res = []
    for i, c in enumerate(row):
        if isinstance(c, tuple):
            res.append(c)
        else:
            res.append((c if c is not None else '', i, i + 1))
    return res


def _chevauche(a, b):
    return a[1] < b[2] and b[1] < a[2]


def _paires_grille(rows):
    """Lit une grille entière.

    (a) en-tête AU-DESSUS de ses valeurs — la بطاقة officielle : on descend
        ligne après ligne sous l'en-tête et on prend TOUT ce qui s'y trouve,
        jusqu'à la prochaine ligne d'en-têtes. Un محور par ligne, dix محاور
        sur dix lignes : les dix sont lus.
    (b) « libellé | valeur » sur la même ligne — le قالب à deux colonnes —,
        avec les lignes de continuation dont la colonne du libellé est vide.
    (c) en-tête et valeurs dans la MÊME case (puces sous le titre).
    """
    grille = [_cases(r) for r in rows]
    nb = len(grille)
    acc = {}

    def ligne_d_entetes(r):
        return any(_est_libelle(c[0]) for c in grille[r] if _nettoyer(c[0]))

    for r in range(nb):
        for idx, case in enumerate(grille[r]):
            ch, reste = _entete(case[0])
            if not ch or ch == IGNORE or ch in acc:
                continue

            valeurs = list(reste)                                   # (c)

            # (a) vers le bas, dans les colonnes de l'en-tête
            for r2 in range(r + 1, nb):
                if ligne_d_entetes(r2):
                    break
                for dessous in grille[r2]:
                    if _chevauche(dessous, case):
                        for l in _lignes(dessous[0]):
                            if l not in valeurs:
                                valeurs.append(l)

            # (b) à côté, sur la même ligne, puis continuation
            if not valeurs and idx + 1 < len(grille[r]):
                voisine = grille[r][idx + 1]
                if not _est_libelle(voisine[0]):
                    valeurs = _lignes(voisine[0])
                    for r2 in range(r + 1, nb):
                        sous_lib = [c for c in grille[r2] if _chevauche(c, case)]
                        if any(_nettoyer(c[0]) for c in sous_lib):
                            break                   # nouveau libellé : fin
                        suite = [c for c in grille[r2] if _chevauche(c, voisine)]
                        lignes = [l for c in suite for l in _lignes(c[0])]
                        if not lignes:
                            continue
                        valeurs.extend(l for l in lignes if l not in valeurs)

            if valeurs:
                acc[ch] = valeurs
    return _finaliser(acc)


# ─── DOCX ────────────────────────────────────────────────────────────────────

_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _texte_paragraphe(p):
    morceaux = []
    for el in p.iter():
        if el.tag == f'{_W}t':
            morceaux.append(el.text or '')
        elif el.tag in (f'{_W}br', f'{_W}cr'):
            morceaux.append('\n')                  # saut de ligne manuel (Maj+Entrée)
        elif el.tag == f'{_W}tab':
            morceaux.append(' ')
    return ''.join(morceaux)


def _texte_cellule(tc):
    paras = [_texte_paragraphe(p) for p in tc.iter(f'{_W}p')]
    return '\n'.join(paras).strip()


def _grilles_docx(root):
    """Chaque tableau → lignes de cases (texte, c0, c1).

    `gridSpan` élargit une case ; une case `vMerge` de continuation (vide)
    est ignorée : sa valeur est déjà celle de la case du dessus. Les tableaux
    imbriqués sont lus pour eux-mêmes."""
    grilles = []
    for tbl in root.iter(f'{_W}tbl'):
        rows = []
        for tr in tbl.findall(f'{_W}tr'):
            ligne, col = [], 0
            for tc in tr.findall(f'{_W}tc'):
                pr = tc.find(f'{_W}tcPr')
                span, continuation = 1, False
                if pr is not None:
                    gs = pr.find(f'{_W}gridSpan')
                    if gs is not None:
                        try:
                            span = max(1, int(gs.get(f'{_W}val', '1')))
                        except ValueError:
                            span = 1
                    vm = pr.find(f'{_W}vMerge')
                    if vm is not None and vm.get(f'{_W}val', 'continue') != 'restart':
                        continuation = True
                texte = '' if continuation else _texte_cellule(tc)
                ligne.append((texte, col, col + span))
                col += span
            rows.append(ligne)
        if rows:
            grilles.append(rows)
    return grilles


def _paras_docx(root):
    """Paragraphes hors tableaux, dans l'ordre du document."""
    body = root.find(f'{_W}body')
    if body is None:
        return []
    return [_texte_paragraphe(el) for el in body if el.tag == f'{_W}p']


def _depuis_paras(cellules, strict=False):
    """Texte libre : un en-tête ouvre une خانة, et TOUTES les lignes qui
    suivent lui appartiennent jusqu'au prochain en-tête (repris ou non).

    strict : le texte provient d'un tableau APLATI (Excel, .doc). Deux
    en-têtes qui se suivent sans rien entre eux y sont deux colonnes voisines
    — «محاور | أهداف» — et les lignes qui suivent mélangent les deux. On ne
    devine pas : aucune de ces colonnes ne reçoit ce texte."""
    acc = {}
    courant = None
    en_tete_precedent = False
    for brut in cellules:
        ch, reste = _entete(brut)
        if ch:
            if strict and en_tete_precedent:
                # colonnes voisines : on retire ce qui vient d'être ouvert
                if courant and not acc.get(courant):
                    acc.pop(courant, None)
                courant = None
                en_tete_precedent = True
                continue
            courant = ch if (ch != IGNORE and ch not in acc) else None
            if courant:
                acc[courant] = []
                _ajouter(acc, courant, reste)
            en_tete_precedent = not reste
            continue
        en_tete_precedent = False
        if courant:
            _ajouter(acc, courant, _lignes(brut))
    return _finaliser(acc)


def _lire_docx(data):
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
        xml = z.read('word/document.xml')
    except Exception:
        _log.warning('_lire_docx : exception ignorée', exc_info=True)
        raise ErreurBataqa('تعذّرت قراءة ملفّ وورد. يُقبل شكل .docx.')
    root = ET.fromstring(xml)
    res = {}
    for grille in _grilles_docx(root):
        _fusion(res, _paires_grille(grille))
    paras = _paras_docx(root)
    _fusion(res, _depuis_paras(paras))
    if not res.get('titre'):
        t = _titre_depuis_texte(paras)
        if t:
            res['titre'] = t
    return res


# ─── XLSX / XLS ──────────────────────────────────────────────────────────────

def _texte_brut(v):
    """Comme `_texte`, mais GARDE les retours à la ligne d'une case : une
    case Excel à plusieurs lignes porte plusieurs محاور."""
    if v is None or isinstance(v, bool):
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _lire_xlsx(data):
    try:
        from openpyxl import load_workbook
    except Exception:
        _log.warning('_lire_xlsx : exception ignorée', exc_info=True)
        raise ErreurBataqa('تعذّر تحميل قارئ ملفّات إكسال (openpyxl).')
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        _log.warning('_lire_xlsx : exception ignorée', exc_info=True)
        raise ErreurBataqa('تعذّرت قراءة الملفّ. يُقبل شكل .xlsx.')
    res = {}
    for ws in wb.worksheets:
        # Cases fusionnées : la case d'origine couvre toute la plage ; les
        # autres cases de la plage n'existent pas pour la lecture.
        couvre, masquees = {}, set()
        for plage in ws.merged_cells.ranges:
            couvre[(plage.min_row, plage.min_col)] = plage.max_col + 1
            for r in range(plage.min_row, plage.max_row + 1):
                for c in range(plage.min_col, plage.max_col + 1):
                    if (r, c) != (plage.min_row, plage.min_col):
                        masquees.add((r, c))
        rows = []
        for r in range(1, (ws.max_row or 0) + 1):
            ligne = []
            for c in range(1, (ws.max_column or 0) + 1):
                if (r, c) in masquees:
                    continue
                fin = couvre.get((r, c), c + 1)
                ligne.append((_texte_brut(ws.cell(r, c).value), c, fin))
            rows.append(ligne)
        _fusion(res, _paires_grille(rows))
        # Une feuille écrite comme un texte (une colonne, en-têtes et lignes)
        _fusion(res, _depuis_paras([c[0] for ligne in rows for c in ligne if c[0]],
                                   strict=True))
    return res


def _lire_xls(data):
    try:
        import xlrd
    except Exception:
        _log.warning('_lire_xls : exception ignorée', exc_info=True)
        raise ErreurBataqa('لقراءة .xls القديم ثبّت المكتبة xlrd، أو احفظ '
                           'الملفّ بصيغة .xlsx.')
    try:
        wb = xlrd.open_workbook(file_contents=data)
    except Exception:
        _log.warning('_lire_xls : exception ignorée', exc_info=True)
        raise ErreurBataqa('تعذّرت قراءة ملفّ .xls.')
    res = {}
    for sh in wb.sheets():
        rows = [[_texte_brut(sh.cell_value(r, c)) for c in range(sh.ncols)]
                for r in range(sh.nrows)]
        _fusion(res, _paires_grille(rows))
        _fusion(res, _depuis_paras([c for ligne in rows for c in ligne if c], strict=True))
    return res


# ─── DOC (Word 97-2003) ──────────────────────────────────────────────────────
#
# Un .doc est un conteneur OLE (« Compound File ») qui renferme le flux
# « WordDocument ». Le texte n'y est pas d'un seul tenant : une table des
# morceaux (« piece table », dans le flux 0Table/1Table) dit où se trouve
# chaque tronçon et s'il est en Unicode ou en 8 bits. On la suit : c'est le
# SEUL moyen d'obtenir le texte du document, dans l'ordre, sans les octets
# internes du conteneur (noms de styles, tables…) qu'un simple balayage
# ramassait aussi.
#
# Dans ce texte, chaque case de tableau finit par \x07, et chaque LIGNE de
# tableau finit par un \x07 de plus : une ligne de k cases donne k+1 marques.
# C'est ce compte qui permet de savoir, sans deviner, quelle case est sous
# quel en-tête — même quand une case est vide ou ne contient qu'un nombre.

_FIN = 0xFFFFFFFE
_LIBRE = 0xFFFFFFFF


class _Conteneur:
    """Lecteur minimal de Compound File Binary (MS-CFB), sans dépendance."""

    def __init__(self, data):
        import struct
        self.s = struct
        if data[:8] != b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
            raise ErreurBataqa('الملفّ ليس ملفّ وورد .doc صالحًا.')
        self.d = data
        u16 = lambda o: struct.unpack_from('<H', data, o)[0]
        u32 = lambda o: struct.unpack_from('<I', data, o)[0]
        self.taille = 1 << u16(0x1E)
        self.mini_taille = 1 << u16(0x20)
        self.seuil = u32(0x38)
        # FAT, par le DIFAT (109 entrées dans l'en-tête, puis secteurs chaînés)
        difat = [u32(0x4C + 4 * i) for i in range(109)]
        sect, reste = u32(0x44), u32(0x48)
        n = self.taille // 4
        while sect < _FIN and reste > 0:
            base = (sect + 1) * self.taille
            difat += [u32(base + 4 * i) for i in range(n - 1)]
            sect = u32(base + 4 * (n - 1))
            reste -= 1
        self.fat = []
        for s in difat:
            if s >= _FIN:
                continue
            base = (s + 1) * self.taille
            self.fat += list(struct.unpack_from(f'<{n}I', data, base))
        # Répertoire
        brut = self._chaine(u32(0x30))
        self.entrees = {}
        racine = None
        for o in range(0, len(brut) - 127, 128):
            lg = struct.unpack_from('<H', brut, o + 64)[0]
            nom = brut[o:o + max(0, lg - 2)].decode('utf-16-le', 'ignore')
            typ = brut[o + 66]
            debut = struct.unpack_from('<I', brut, o + 116)[0]
            taille = struct.unpack_from('<I', brut, o + 120)[0]
            if typ == 5:
                racine = (debut, taille)
            if typ in (2, 5):
                self.entrees[nom] = (debut, taille)
        self.mini_flux = self._chaine(racine[0])[:racine[1]] if racine else b''
        self.mini_fat = []
        mf = u32(0x3C)
        if mf < _FIN:
            m = self._chaine(mf)
            self.mini_fat = list(struct.unpack_from(f'<{len(m) // 4}I', m, 0))

    def _chaine(self, sect):
        morceaux, vus = [], set()
        while sect < _FIN and sect not in vus and sect < len(self.fat) + 1:
            vus.add(sect)
            base = (sect + 1) * self.taille
            morceaux.append(self.d[base:base + self.taille])
            if sect >= len(self.fat):
                break
            sect = self.fat[sect]
        return b''.join(morceaux)

    def flux(self, nom):
        if nom not in self.entrees:
            return None
        debut, taille = self.entrees[nom]
        if taille < self.seuil:
            morceaux, sect, vus = [], debut, set()
            while sect < _FIN and sect not in vus and sect < len(self.mini_fat):
                vus.add(sect)
                o = sect * self.mini_taille
                morceaux.append(self.mini_flux[o:o + self.mini_taille])
                sect = self.mini_fat[sect]
            return b''.join(morceaux)[:taille]
        return self._chaine(debut)[:taille]


def _texte_word97(data):
    """Le texte principal d'un .doc, tel que Word le stocke (\r, \x07…)."""
    st = __import__('struct')
    c = _Conteneur(data)
    wd = c.flux('WordDocument')
    if not wd or len(wd) < 0x200 or st.unpack_from('<H', wd, 0)[0] != 0xA5EC:
        raise ErreurBataqa('الملفّ ليس وثيقة وورد .doc.')
    drapeaux = st.unpack_from('<H', wd, 0x0A)[0]
    if drapeaux & 0x0100:
        raise ErreurBataqa('الملفّ محميّ بكلمة سرّ: احفظه دون حماية ثمّ أعد المحاولة.')
    table = c.flux('1Table' if drapeaux & 0x0200 else '0Table')
    if not table:
        raise ErreurBataqa('تعذّرت قراءة بنية ملفّ .doc.')
    csw = st.unpack_from('<H', wd, 32)[0]
    o = 34 + 2 * csw
    cslw = st.unpack_from('<H', wd, o)[0]
    lw = o + 2
    ccp_text = st.unpack_from('<I', wd, lw + 12)[0]
    o = lw + 4 * cslw
    rg = o + 2
    fc_clx, lcb_clx = st.unpack_from('<II', wd, rg + 33 * 8)
    clx = table[fc_clx:fc_clx + lcb_clx]
    i = 0
    while i < len(clx) and clx[i] == 0x01:                  # Prc : on saute
        i += 3 + st.unpack_from('<h', clx, i + 1)[0]
    if i >= len(clx) or clx[i] != 0x02:
        raise ErreurBataqa('تعذّرت قراءة بنية ملفّ .doc.')
    lcb = st.unpack_from('<I', clx, i + 1)[0]
    plc = clx[i + 5:i + 5 + lcb]
    n = (lcb - 4) // 12
    cps = st.unpack_from(f'<{n + 1}I', plc, 0)
    texte = []
    for k in range(n):
        debut, fin = cps[k], min(cps[k + 1], ccp_text)
        if debut >= fin:
            continue
        fc = st.unpack_from('<I', plc, 4 * (n + 1) + 8 * k + 2)[0]
        if fc & 0x40000000:                                  # 8 bits (cp1252)
            fc = (fc & 0x3FFFFFFF) // 2
            texte.append(wd[fc:fc + (fin - debut)].decode('cp1252', 'replace'))
        else:
            texte.append(wd[fc:fc + 2 * (fin - debut)].decode('utf-16-le', 'replace'))
    return ''.join(texte)


def _nettoyer_word(texte):
    """Retire les codes de champ (\x13 instruction \x14 résultat \x15 : on
    garde le résultat) et les caractères de contrôle d'objets."""
    res, pile = [], []
    for ch in texte:
        if ch == '\x13':
            pile.append('instr'); continue
        if ch == '\x14':
            if pile: pile[-1] = 'res'
            continue
        if ch == '\x15':
            if pile: pile.pop()
            continue
        if pile and pile[-1] == 'instr':
            continue
        if ch == '\x0b':
            res.append('\r')                                 # saut de ligne manuel
        elif ch in '\x01\x08\x0c\x02\x03\x04\x05':
            continue
        elif ch in '\x1e\x1f':
            res.append('-' if ch == '\x1e' else '')
        else:
            res.append(ch)
    return ''.join(res)


def _elements_doc(texte, est_entete=None):
    """Le texte d'un .doc → suite d'éléments ('hors', lignes) / ('case', texte).

    Une case = ce qui précède un \x07. Mais le texte HORS tableau qui précède
    la première case d'un tableau (titre du document, titre de section…) lui
    est collé, sans séparateur : on l'en détache au premier EN-TÊTE de
    colonne rencontré. Le dernier morceau (après la dernière marque) est
    hors tableau."""
    morceaux = texte.split('\x07')
    elements = []
    for idx, m in enumerate(morceaux):
        lignes = [l for l in m.split('\r')]
        dernier = idx == len(morceaux) - 1
        if dernier:
            if any(_nettoyer(l) for l in lignes):
                elements.append(('hors', lignes))
            break
        # Détacher le texte hors tableau collé devant une case d'en-tête.
        coupe = None
        reconnait = est_entete or _classer_ligne
        for j, l in enumerate(lignes):
            if not _nettoyer(l) or _est_section(l):
                continue
            if reconnait(l):
                coupe = j
                break
        if coupe:
            avant = lignes[:coupe]
            if any(_nettoyer(l) for l in avant):
                elements.append(('hors', avant))
            lignes = lignes[coupe:]
        elements.append(('case', '\n'.join(l for l in lignes)))
    return elements


def _vide(e):
    return e[0] == 'case' and not _nettoyer(e[1].replace('\n', ' '))


def _depuis_doc(elements):
    """Lit les tableaux d'un .doc à partir de ses éléments.

    (a) Ligne d'EN-TÊTES suivie de sa marque de fin de ligne : tableau à k
        colonnes. Chaque ligne suivante compte k cases puis une marque ; la
        case i va à l'en-tête i. Si une ligne ne respecte pas ce compte (cases
        fusionnées…), on s'arrête là pour ce tableau plutôt que de décaler.
    (b) « libellé | valeur » sur la même ligne (قالب, fiches) : chaque
        libellé prend les cases qui le suivent jusqu'au libellé suivant.
    """
    acc = {}
    i, n = 0, len(elements)

    def entete(e):
        return e[0] == 'case' and _entete(e[1])[0] is not None

    def champ(e):
        return _entete(e[1])[0]

    def verser(ch, texte):
        if ch and ch != IGNORE:
            if ch not in acc:
                acc[ch] = []
            _ajouter(acc, ch, _lignes(texte))

    while i < n:
        e = elements[i]
        if not entete(e):
            i += 1
            continue
        # En-têtes consécutifs de la même ligne
        j = i
        entetes = []
        while j < n and entete(elements[j]):
            ch, reste = _entete(elements[j][1])
            entetes.append(ch if (ch != IGNORE and ch not in acc) else IGNORE)
            if ch != IGNORE and ch not in acc:
                acc[ch] = []
                _ajouter(acc, ch, reste)
            j += 1
        if j < n and _vide(elements[j]):
            # (a) tableau vertical à k colonnes
            k = len(entetes)
            j += 1
            while j + k < n + 1:
                ligne = elements[j:j + k]
                if len(ligne) < k or any(x[0] != 'case' or entete(x) for x in ligne):
                    break
                marque = elements[j + k] if j + k < n else None
                if marque is None or not _vide(marque):
                    break                         # compte rompu : on n'invente rien
                for col, x in enumerate(ligne):
                    verser(entetes[col], x[1])
                j += k + 1
            i = j
            continue
        # (b) libellé | valeur sur la même ligne
        courant = entetes[-1]
        while j < n and elements[j][0] == 'case' and not _vide(elements[j]):
            x = elements[j]
            if entete(x):
                ch, reste = _entete(x[1])
                courant = ch if (ch != IGNORE and ch not in acc) else IGNORE
                if courant != IGNORE:
                    acc[courant] = []
                    _ajouter(acc, courant, reste)
            else:
                verser(courant, x[1])
            j += 1
        i = j
    return _finaliser(acc)


#: « الدورة التكوينيّة حول " X " » / « دورة تكوينيّة في مجال " X " » : le
#: عنوان figure souvent aussi dans le titre du document.
_RE_TITRE = re.compile(r'(?:حول|في\s+مجال)\s*["«“”]\s*(.+?)\s*["»”“]')


def _titre_depuis_texte(lignes):
    for l in lignes:
        m = _RE_TITRE.search(_nettoyer(l))
        if m and 2 < len(m.group(1)) <= 120:
            return m.group(1).strip()
    return ''


def _lire_doc(data):
    try:
        texte = _nettoyer_word(_texte_word97(data))
    except ErreurBataqa:
        raise
    except Exception:
        _log.warning('_lire_doc : lecture structurée impossible', exc_info=True)
        raise ErreurBataqa('تعذّرت القراءة الدقيقة لملفّ .doc القديم. '
                           'يُرجى حفظه بصيغة .docx ثمّ إعادة المحاولة.')
    elements = _elements_doc(texte)
    res = _depuis_doc(elements)
    # Texte hors tableau : « libellé : valeur » libre, jamais au-dessus du tableau.
    hors = [l for e in elements if e[0] == 'hors' for l in e[1]]
    _fusion(res, _depuis_paras(hors))
    if not res.get('titre'):
        t = _titre_depuis_texte(hors or texte.split('\r'))
        if t:
            res['titre'] = t
    if not res:
        raise ErreurBataqa('لم يتعرّف النظام على بطاقة بيداغوجيّة في ملفّ .doc. '
                           'يُرجى حفظه بصيغة .docx ثمّ إعادة المحاولة.')
    return res


# ─── Entrée ──────────────────────────────────────────────────────────────────

def analyser(data, filename):
    """data : octets ; filename : pour l'extension. Renvoie {خانة: valeur}."""
    ext = os.path.splitext(filename or '')[1].lower()
    if ext in ('.xlsx', '.xlsm'):
        res = _lire_xlsx(data)
    elif ext == '.xls':
        res = _lire_xls(data)
    elif ext == '.docx':
        res = _lire_docx(data)
    elif ext == '.doc':
        res = _lire_doc(data)
    else:
        raise ErreurBataqa('صيغة غير مدعومة. حمّل .xlsx أو .xls أو .docx أو .doc '
                           '(أو استعمل القالب الجاهز).')
    # Le نوع doit tomber dans la liste fermée du formulaire. « تكوين مستمرّ
    # حضوريّ » ou « مستمر (حضوري) » y sont ramenés ; ce qui n'y ressemble pas
    # est ignoré plutôt que deviné.
    tf = res.get('type_formation', '')
    if tf:
        n = _normaliser(tf)
        if 'مستمر' in n and 'حضوري' in n:
            res['type_formation'] = 'مستمر حضوري'
        elif 'مستمر' in n:
            res['type_formation'] = 'مستمر'
        elif 'اساسي' in n:
            res['type_formation'] = 'أساسي'
        else:
            res.pop('type_formation', None)
    # Les خانات à une seule ligne dans le formulaire ne gardent que la première.
    for cle in ('titre', 'lieu_formation_defaut'):
        if res.get(cle):
            res[cle] = res[cle].splitlines()[0].strip()
    return res


def construire_modele():
    """Le قالب .xlsx à télécharger : deux colonnes (البيان | القيمة), une ligne
    par خانة — relu à 100% par la منظومة."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = 'بطاقة'
    try:
        ws.sheet_view.rightToLeft = True
    except Exception:
        pass
    ws['A1'] = 'البيان'
    ws['B1'] = 'القيمة'
    fill = PatternFill('solid', fgColor='5B9BD5')
    for c in ('A1', 'B1'):
        ws[c].font = Font(bold=True, color='FFFFFF')
        ws[c].fill = fill
        ws[c].alignment = Alignment(horizontal='center')
    lignes = [
        ('عنوان التكوين', ''), ('نوع التكوين', 'أساسي'),
        ('المستهدفون بالتكوين', ''), ('مكان التكوين', ''),
        ('محاور الدورة', ''), ('أهداف الدورة', ''),
        ('الطرق البيداغوجية', ''), ('المعينات البيداغوجية', ''),
        ('الإعداد المادي', ''), ('التجهيزات والمعدات الخصوصية', ''),
    ]
    for i, (lab, val) in enumerate(lignes, start=2):
        ws[f'A{i}'] = lab
        ws[f'B{i}'] = val
        ws[f'A{i}'].font = Font(bold=True)
        ws[f'A{i}'].alignment = Alignment(horizontal='right')
        ws[f'B{i}'].alignment = Alignment(horizontal='right', wrap_text=True)
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 55
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
