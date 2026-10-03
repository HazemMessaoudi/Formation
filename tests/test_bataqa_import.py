# -*- coding: utf-8 -*-
"""تحميل دورة — lecture d'une بطاقة (Excel/Word/قالب) pour pré-remplir une مادّة.

On défend : le قالب que la منظومة produit se relit à 100% ; un .docx quelconque
est lu sans dépendance ; rien n'est deviné (libellé inconnu ignoré) ; le route
signale un عنوان déjà présent (existe_id) ; et « تعويض » met à jour la fiche
existante au lieu de créer un doublon."""

import io
import zipfile

from openpyxl import load_workbook

from core import bataqa_import as B


JETON = {'X-CSRF-Token': 'jeton-de-test'}


def _modele_rempli(valeurs):
    data = B.construire_modele()
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    for r in range(2, ws.max_row + 1):
        lab = ws.cell(r, 1).value
        if lab in valeurs:
            ws.cell(r, 2).value = valeurs[lab]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx(paragraphes):
    NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    body = ''.join(
        f'<w:p><w:r><w:t xml:space="preserve">{t}</w:t></w:r></w:p>'
        for t in paragraphes)
    doc = (f'<?xml version="1.0"?><w:document xmlns:w="{NS}">'
           f'<w:body>{body}</w:body></w:document>')
    b = io.BytesIO()
    with zipfile.ZipFile(b, 'w') as z:
        z.writestr('word/document.xml', doc)
    return b.getvalue()


# ── unité : le lecteur ───────────────────────────────────────────────────────

def test_qalab_se_relit_a_100():
    data = _modele_rempli({
        'عنوان التكوين': 'إجراءات التبليغ',
        'المستهدفون بالتكوين': 'أعوان الديوانة',
        'محاور الدورة': 'المحور الأوّل',
        'أهداف الدورة': 'إتقان الإجراء'})
    r = B.analyser(data, 'q.xlsx')
    assert r['titre'] == 'إجراءات التبليغ'
    assert r['mustahdafun'] == 'أعوان الديوانة'
    assert r['mahawer'] == 'المحور الأوّل'
    assert r['type_formation'] == 'أساسي'


def test_docx_lu_sans_dependance():
    data = _docx(['عنوان التكوين', 'مكافحة التهريب', 'أهداف الدورة', 'ردع المخالفات'])
    r = B.analyser(data, 'd.docx')
    assert r['titre'] == 'مكافحة التهريب'
    assert r['objectifs'] == 'ردع المخالفات'


def _docx_tables(tables):
    """tables : liste de tableaux ; chaque tableau = liste de lignes ; chaque
    case peut contenir des retours à la ligne (puces)."""
    NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

    def p(t):
        return f'<w:p><w:r><w:t xml:space="preserve">{t}</w:t></w:r></w:p>'

    def tc(cell):
        return '<w:tc>' + ''.join(p(x) for x in cell.split('\n')) + '</w:tc>'

    def tr(cells):
        return '<w:tr>' + ''.join(tc(c) for c in cells) + '</w:tr>'

    def tbl(rows):
        return '<w:tbl>' + ''.join(tr(r) for r in rows) + '</w:tbl>'

    body = ''.join(tbl(t) for t in tables)
    doc = (f'<?xml version="1.0"?><w:document xmlns:w="{NS}">'
           f'<w:body>{body}</w:body></w:document>')
    b = io.BytesIO()
    with zipfile.ZipFile(b, 'w') as z:
        z.writestr('word/document.xml', doc)
    return b.getvalue()


def test_bataqa_reelle_en_tableaux():
    """La بطاقة réelle est une grille en-tête/valeur (le même modèle que la
    منظومة imprime). C'est le cas qui échouait : on lit la valeur SOUS chaque
    en-tête, et les listes à puces d'une case sont recollées entières."""
    mahawer = ('- أصناف الأسلحة المتوفرة وخصائصها\n'
               '- إستعمال و تفكيك و صيانة الأسلحة المعتمدة\n'
               '- دروس نظرية حول تحوطات الأمان')
    ahdaf = ('- التعرّف على أصناف الأسلحة والذّخيرة\n'
             '- التمكّن من تحوّطات الأمان.')
    data = _docx_tables([
        [['نوع التكوين', 'عدد المشاركين', 'المستهدفون بالتكوين', 'موضوع التكوين'],
         ['مستمر حضوري', '12', 'ضبّاط الديوانة', 'إستعمال وصيانة الأسلحة']],
        [['أهداف الدورة', 'محاور الدورة'], [ahdaf, mahawer]],
        [['تاريخ الدورة التكوينية', 'مكان التكوين'],
         ['26 فيفري 2026', 'مركز التكوين الديواني بالوسط الغربي بالقصرين']],
    ])
    r = B.analyser(data, 'بطاقة.docx')
    assert r['titre'] == 'إستعمال وصيانة الأسلحة'
    assert r['type_formation'] == 'مستمر حضوري'
    assert r['mustahdafun'] == 'ضبّاط الديوانة'
    assert r['lieu_formation_defaut'] == 'مركز التكوين الديواني بالوسط الغربي بالقصرين'
    assert 'دروس نظرية حول تحوطات الأمان' in r['mahawer']
    assert 'التمكّن من تحوّطات الأمان' in r['objectifs']
    # les puces « - » de début de ligne sont retirées
    assert not r['mahawer'].lstrip().startswith('-')


def test_rien_nest_devine():
    # Un libellé inconnu ne doit remplir aucune خانة.
    data = _docx(['حقل غير معروف تمامًا', 'قيمة ما'])
    assert B.analyser(data, 'd.docx') == {}


def test_format_refuse():
    import pytest
    with pytest.raises(B.ErreurBataqa):
        B.analyser(b'xx', 'photo.png')


# ── routes ───────────────────────────────────────────────────────────────────

def test_route_modele_telecharge_un_xlsx(client):
    r = client.get('/mawad/modele-bataqa')
    assert r.status_code == 200
    assert r.data[:2] == b'PK'          # un .xlsx est une archive ZIP


def test_route_importer_signale_le_titre_existant(client, db):
    db.add_madda({'titre': 'إجراءات التبليغ', 'type_formation': 'أساسي'})
    data = _modele_rempli({'عنوان التكوين': 'إجراءات التبليغ',
                           'أهداف الدورة': 'الإتقان'})
    r = client.post('/mawad/importer', headers=JETON,
                    data={'fichier': (io.BytesIO(data), 'q.xlsx')},
                    content_type='multipart/form-data')
    j = r.get_json()
    assert j['succes'] is True
    assert j['champs']['titre'] == 'إجراءات التبليغ'
    assert j['existe_id']       # le titre existe déjà


def test_remplacer_ne_cree_pas_de_doublon(client, db):
    db.add_madda({'titre': 'مكافحة التهريب', 'type_formation': 'أساسي',
                  'objectifs': 'قديم'})
    ancien = db.get_madda_by_titre('مكافحة التهريب')['id']
    avant = len(db.get_mawad())
    r = client.post('/mawad/ajouter', headers=JETON, data={
        '_csrf': 'jeton-de-test', 'remplacer_id': str(ancien),
        'titre': 'مكافحة التهريب', 'type_formation': 'مستمر',
        'objectifs': 'جديد'}, follow_redirects=False)
    assert r.status_code in (302, 200)
    apres = db.get_mawad()
    assert len(apres) == avant                 # pas de doublon
    maj = db.get_madda(ancien)
    assert maj['objectifs'] == 'جديد'          # la fiche a bien été remplacée
    assert maj['type_formation'] == 'مستمر'


# ── v1.4d : lire TOUT le fichier ─────────────────────────────────────────────
#
# Incident : pour une بطاقة réelle, «محاور الدورة» et «أهداف الدورة» ne
# recevaient que leur première ligne ; «المستهدفون بالتكوين» recevait
# «عدد المشاركين» et «مكان التكوين» recevait «تاريخ الدّورة التّكوينيّة» —
# des EN-TÊTES voisins pris pour des valeurs parce que la case était vide.

MAHAWER = ['تقديم الإطار العامّ للدّرس', 'مفهوم الفاتورة وأنواعها',
           'البيانات الوجوبيّة للفاتورة', 'الفاتورة الإلكترونيّة',
           'المخالفات المتعلّقة بالفوترة']
AHDAF = ['التعرّف على الإطار القانوني للفوترة', 'التمييز بين أنواع الفواتير',
         'التثبّت من صحّة الفاتورة عند المراقبة']


def _docx_grille(tables):
    """Comme _docx_tables, mais une case peut être (texte, {'span': n}) ou
    (texte, {'vmerge': 'restart'|'continue'}) ; et « \\n » dans un texte
    devient un saut de ligne manuel <w:br/> quand la case le demande."""
    NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

    def p(t, br=False):
        if br:
            runs = '<w:br/>'.join(f'<w:t xml:space="preserve">{x}</w:t>' for x in t.split('\n'))
            return f'<w:p><w:r>{runs}</w:r></w:p>'
        return f'<w:p><w:r><w:t xml:space="preserve">{t}</w:t></w:r></w:p>'

    def tc(cell):
        opts = {}
        if isinstance(cell, tuple):
            cell, opts = cell
        pr = ''
        if opts.get('span'):
            pr += f'<w:gridSpan w:val="{opts["span"]}"/>'
        if opts.get('vmerge') == 'restart':
            pr += '<w:vMerge w:val="restart"/>'
        elif opts.get('vmerge') == 'continue':
            pr += '<w:vMerge/>'
        pr = f'<w:tcPr>{pr}</w:tcPr>' if pr else ''
        if opts.get('br'):
            corps = p(cell, br=True)
        else:
            corps = ''.join(p(x) for x in (cell.split('\n') if cell else ['']))
        return f'<w:tc>{pr}{corps}</w:tc>'

    body = ''.join('<w:tbl>' + ''.join('<w:tr>' + ''.join(tc(c) for c in r) + '</w:tr>'
                                       for r in t) + '</w:tbl>' for t in tables)
    doc = (f'<?xml version="1.0"?><w:document xmlns:w="{NS}">'
           f'<w:body>{body}</w:body></w:document>')
    b = io.BytesIO()
    with zipfile.ZipFile(b, 'w') as z:
        z.writestr('word/document.xml', doc)
    return b.getvalue()


def _bataqa_officielle_un_axe_par_ligne():
    """La grille imprimée par la منظومة, ordre logique Word (droite → gauche),
    cases vides là où le modèle les laisse vides, un محور/هدف par ligne."""
    lignes_axes = [[MAHAWER[i], AHDAF[i] if i < len(AHDAF) else '']
                   for i in range(len(MAHAWER))]
    return _docx_grille([
        [['موضوع التكوين', 'المستهدفون بالتكوين', 'عدد المشاركين',
          'المصالح المعنية بالمشاركة', 'نوع التكوين'],
         ['الفوترة', '', '', '', 'تكوين مستمرّ حضوريّ']],
        [['محاور الدورة', 'أهداف الدورة']] + lignes_axes,
        [['مكان التكوين', 'تاريخ الدورة التكوينية'], ['', '']],
        [['الطرق البيداغوجية', 'المعينات البيداغوجية'],
         ['العرض النظري', 'حاسوب'], ['دراسة حالات', 'جهاز عرض']],
    ])


def test_tous_les_محاور_et_أهداف_sont_lus():
    r = B.analyser(_bataqa_officielle_un_axe_par_ligne(), 'بطاقة.docx')
    assert r['mahawer'].splitlines() == MAHAWER
    assert r['objectifs'].splitlines() == AHDAF
    assert r['methodes_pedagogiques'].splitlines() == ['العرض النظري', 'دراسة حالات']
    assert r['moyens_pedagogiques'].splitlines() == ['حاسوب', 'جهاز عرض']


def test_un_en_tete_voisin_n_est_jamais_pris_pour_une_valeur():
    r = B.analyser(_bataqa_officielle_un_axe_par_ligne(), 'بطاقة.docx')
    assert 'mustahdafun' not in r, f"المستهدفون ← «{r.get('mustahdafun')}»"
    assert 'lieu_formation_defaut' not in r, f"المكان ← «{r.get('lieu_formation_defaut')}»"
    for v in r.values():
        assert 'عدد المشاركين' not in v and 'تاريخ الدورة' not in v


def test_un_contenu_qui_ressemble_a_un_en_tete_reste_du_contenu():
    """« تقديم الإطار العامّ للدّرس » contient les mots d'un en-tête ignoré :
    c'est pourtant le premier محور, et la lecture ne doit pas s'y arrêter."""
    r = B.analyser(_bataqa_officielle_un_axe_par_ligne(), 'بطاقة.docx')
    assert r['mahawer'].splitlines()[0] == 'تقديم الإطار العامّ للدّرس'
    assert len(r['mahawer'].splitlines()) == len(MAHAWER)
    assert B._champ('التعرّف على الأهداف الكبرى للإصلاح الجبائي') is None


def test_le_type_est_ramene_a_la_liste_du_formulaire():
    r = B.analyser(_bataqa_officielle_un_axe_par_ligne(), 'بطاقة.docx')
    assert r['type_formation'] == 'مستمر حضوري'


def test_en_tete_et_puces_dans_la_meme_case():
    data = _docx_grille([[['محاور الدورة\n- ' + '\n- '.join(MAHAWER)],
                          ['أهداف الدورة:\n1- ' + '\n2- '.join(AHDAF)]]])
    r = B.analyser(data, 'x.docx')
    assert r['mahawer'].splitlines() == MAHAWER
    assert r['objectifs'].splitlines() == AHDAF


def test_sauts_de_ligne_manuels_dans_une_case():
    data = _docx_grille([[['محاور الدورة'], [('\n'.join(MAHAWER), {'br': True})]]])
    assert B.analyser(data, 'x.docx')['mahawer'].splitlines() == MAHAWER


def test_en_tete_fusionne_sur_deux_colonnes():
    """Un en-tête étalé sur deux colonnes récupère les deux cases dessous."""
    data = _docx_grille([[[('محاور الدورة', {'span': 2})],
                          [MAHAWER[0], MAHAWER[1]],
                          [MAHAWER[2], MAHAWER[3]]]])
    assert B.analyser(data, 'x.docx')['mahawer'].splitlines() == MAHAWER[:4]


def test_texte_libre_toutes_les_lignes_jusqu_au_titre_suivant():
    data = _docx(['موضوع التكوين: الفوترة', 'محاور الدورة'] + MAHAWER
                 + ['أهداف الدورة'] + AHDAF + ['ملاحظات', 'لا شيء'])
    r = B.analyser(data, 'x.docx')
    assert r['titre'] == 'الفوترة'
    assert r['mahawer'].splitlines() == MAHAWER
    assert r['objectifs'].splitlines() == AHDAF          # « ملاحظات » l'arrête


def test_qalab_valeur_sur_plusieurs_lignes_et_continuation():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(['البيان', 'القيمة'])
    ws.append(['عنوان التكوين', 'الفوترة'])
    ws.append(['محاور الدورة', '\n'.join(MAHAWER[:3])])    # case à plusieurs lignes
    ws.append([None, MAHAWER[3]])                           # continuation
    ws.append([None, MAHAWER[4]])
    ws.append(['أهداف الدورة', AHDAF[0]])
    buf = io.BytesIO()
    wb.save(buf)
    r = B.analyser(buf.getvalue(), 'q.xlsx')
    assert r['mahawer'].splitlines() == MAHAWER
    assert r['objectifs'].splitlines() == [AHDAF[0]]


def test_doc_ancien_repartit_les_cases_par_colonne():
    """Un .doc livre ses cases à plat : chaque case finit par \\x07, et chaque
    LIGNE par un \\x07 de plus. On compte : une case vide ou un simple nombre
    ne décale plus rien (v1.4e)."""
    texte = ''
    texte += 'محاور الدورة\x07أهداف الدورة\x07\x07'
    for i in range(3):
        texte += f'{MAHAWER[i]}\x07{AHDAF[i]}\x07\x07'
    texte += 'مكان التكوين\x07عدد المشاركين\x07\x07'
    texte += '\x0712\x07\x07'                       # lieu VIDE, puis un nombre seul
    r = B._depuis_doc(B._elements_doc(texte + 'fin'))
    assert r['mahawer'].splitlines() == MAHAWER[:3]
    assert r['objectifs'].splitlines() == AHDAF[:3]
    assert 'lieu_formation_defaut' not in r              # vide : rien d'inventé
