# -*- coding: utf-8 -*-
"""توريد المكوّنين من ملفّ إكسال.

Deux niveaux d'épreuve, et ils ne se confondent pas :

* `core.importation` — la lecture et le plan. Rien n'y touche à la قاعدة,
  donc tout s'y éprouve à la ligne près, sans Flask ni base.
* `routes.importation` — la page. On y vérifie surtout ce qui compte pour
  l'agent : que la تجربة n'écrit rien, et que le versement, lui, écrit.

La règle que ces tests défendent avant toute autre : **une بطاقة déjà en
قائمة n'est jamais réécrite par un fichier.** Un tableau d'إكسال est un
apport, jamais une autorité.
"""

from io import BytesIO

import pytest

from core import importation as imp


# ─── outillage ───────────────────────────────────────────────────────────────

def classeur(lignes, titre='المكوّنون'):
    """Un .xlsx en mémoire, construit ligne à ligne."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = titre
    for ligne in lignes:
        ws.append(list(ligne))
    flux = BytesIO()
    wb.save(flux)
    flux.seek(0)
    return flux


ENTETE = ['المعرّف الوحيد', 'الرتبة', 'الاسم', 'اللقب', 'مكان العمل',
          'الجهة المرجعية']


def fiche_db(mkow_id, ident, nom='زياد', prenom='البوهلالي', **reste):
    """Une بطاقة telle que `get_mkowin()` la rend."""
    base = {'id': mkow_id, 'identifiant_unique': ident,
            'nom': nom, 'prenom': prenom, 'grade': '', 'lieu_travail': '',
            'jiha_marjiiya': '', 'telephone_gsm': ''}
    base.update(reste)
    return base


# ─── التّعرّف على العناوين ─────────────────────────────────────────────────────

def test_libelles_reconnus_malgre_les_حركات_والهمزات():
    assert imp.champ_de_libelle('المعرّف الوحيد') == 'identifiant_unique'
    assert imp.champ_de_libelle('المعرف الوحيد') == 'identifiant_unique'
    assert imp.champ_de_libelle('  الرّتبة :') == 'grade'
    assert imp.champ_de_libelle('الإسم') == 'nom'
    assert imp.champ_de_libelle('اللّقب') == 'prenom'
    assert imp.champ_de_libelle('الجهة المرجعية') == 'jiha_marjiiya'
    assert imp.champ_de_libelle('عمود لا يعني شيئا') is None


def test_lordre_du_nom_est_celui_de_la_manzouma():
    """الاسم واللقب — dans cet ordre, et ce test est là pour le verrouiller.

    Partout dans la منظومة le nom affiché est «الاسم واللقب», construit comme
    `nom + ' ' + prenom`. Inverser ces deux خانات inverserait tous les noms
    importés sans qu'aucun autre test ne s'en aperçoive : chaque fiche
    resterait « valide », simplement fausse. D'où cette épreuve explicite.
    """
    assert imp.champ_de_libelle('الاسم') == 'nom'
    assert imp.champ_de_libelle('اللقب') == 'prenom'
    assert imp.COLONNES_MODELE.index('الاسم') < imp.COLONNES_MODELE.index('اللقب')

    plan = imp.analyser(classeur([
        ['المعرّف الوحيد', 'الاسم', 'اللقب'],
        ['77001', 'زياد', 'البوهلالي'],
    ]), [])
    _, fiche = plan['a_ajouter'][0]
    assert fiche['nom'] == 'زياد'
    assert fiche['prenom'] == 'البوهلالي'
    # Tel que la منظومة l'affichera :
    assert f"{fiche['nom']} {fiche['prenom']}" == 'زياد البوهلالي'


def test_entete_cherchee_au_dela_de_la_premiere_ligne():
    """Un tableau administratif porte un titre avant son entête : on cherche."""
    plan = imp.analyser(classeur([
        ['مركز التّكوين الجهوي بالقصرين'],
        [],
        ENTETE,
        ['77001', 'عريف', 'زياد', 'البوهلالي', 'ميناء بنزرت', ''],
    ]), [])
    assert plan['ligne_entete'] == 3
    assert plan['total_lignes'] == 1
    assert len(plan['a_ajouter']) == 1


def test_colonne_inconnue_ignoree_et_signalee():
    """On ignore ce qu'on ne comprend pas — et on le dit, sans bloquer."""
    plan = imp.analyser(classeur([
        ENTETE + ['ملاحظات المدير'],
        ['77001', 'عريف', 'زياد', 'البوهلالي', 'ميناء بنزرت', '', 'س'],
    ]), [])
    assert plan['colonnes_ignorees'] == ['ملاحظات المدير']
    assert 'identifiant_unique' in plan['colonnes_reconnues']
    numero, fiche = plan['a_ajouter'][0]
    assert 'ملاحظات المدير' not in fiche
    assert set(fiche) <= set(c for c, _ in imp.COLONNES)


def test_fichier_sans_entete_reconnaissable_refuse():
    with pytest.raises(imp.ErreurImportation):
        imp.analyser(classeur([['أ', 'ب'], ['ج', 'د']]), [])


def test_une_seule_colonne_reconnue_ne_suffit_pas():
    """Une colonne isolée peut être une coïncidence : on exige deux."""
    with pytest.raises(imp.ErreurImportation):
        imp.analyser(classeur([['الرتبة', 'شيء آخر'], ['عريف', 'س']]), [])


# ─── قراءة الخانات ───────────────────────────────────────────────────────────

def test_identifiant_numerique_ne_devient_pas_flottant():
    """إكسال rend 77001 en flottant ; « 77001.0 » ne correspondrait à rien."""
    plan = imp.analyser(classeur([
        ENTETE,
        [77001, 'عريف', 'زياد', 'البوهلالي', '', ''],
        [77002.0, 'رقيب', 'سمير', 'قاسمي', '', ''],
    ]), [])
    identifiants = [f['identifiant_unique'] for _, f in plan['a_ajouter']]
    assert identifiants == ['77001', '77002']


def test_ligne_vide_ni_acceptee_ni_refusee():
    plan = imp.analyser(classeur([
        ENTETE,
        ['77001', 'عريف', 'زياد', 'البوهلالي', '', ''],
        [None, None, None, None, None, None],
        ['', '', '', '', '', ''],
        ['77002', 'رقيب', 'سمير', 'قاسمي', '', ''],
    ]), [])
    assert plan['total_lignes'] == 2
    assert len(plan['a_ajouter']) == 2
    assert plan['rejetes'] == []


def test_nom_en_une_colonne_jamais_decoupe():
    """« بن عبد الله محمد الهادي » : le système ne tranche pas où est le لقب."""
    plan = imp.analyser(classeur([
        ['اللقب والاسم', 'المعرّف الوحيد'],
        ['بن عبد الله محمد الهادي', '77009'],
    ]), [])
    _, fiche = plan['a_ajouter'][0]
    assert fiche['nom'] == 'بن عبد الله محمد الهادي'
    assert 'prenom' not in fiche
    assert 'nom_complet' not in fiche


def test_le_texte_est_rendu_caractere_pour_caractere():
    """La normalisation sert à reconnaître, jamais à réécrire ce qui est stocké."""
    plan = imp.analyser(classeur([
        ENTETE,
        ['77001', 'عريف', 'زياد', 'البوهلالي', 'ميناء بنزرت',
         'الإدارة الجهويّة للدّيوانة ببنزرت'],
    ]), [])
    _, fiche = plan['a_ajouter'][0]
    assert fiche['jiha_marjiiya'] == 'الإدارة الجهويّة للدّيوانة ببنزرت'
    assert fiche['nom'] == 'زياد'


def test_numeros_de_ligne_sont_ceux_quaffiche_excel():
    plan = imp.analyser(classeur([
        ['عنوان'],
        ENTETE,
        ['77001', 'عريف', 'زياد', 'البوهلالي', '', ''],
        ['', 'عريف', 'فلان', 'بلا معرّف', '', ''],
    ]), [])
    assert plan['a_ajouter'][0][0] == 3
    assert plan['rejetes'][0]['ligne'] == 4


# ─── المعرّف الوحيد مفتاح ─────────────────────────────────────────────────────

def test_ligne_sans_nom_refusee():
    plan = imp.analyser(classeur([
        ENTETE,
        ['77001', 'عريف', '', '', 'ميناء بنزرت', ''],
    ]), [])
    assert plan['a_ajouter'] == []
    assert plan['rejetes'][0]['motif'] == 'بدون اسم'


def test_ligne_sans_identifiant_refusee_par_defaut():
    plan = imp.analyser(classeur([
        ENTETE,
        ['', 'عريف', 'زياد', 'البوهلالي', '', ''],
    ]), [])
    assert plan['a_ajouter'] == []
    assert plan['rejetes'][0]['motif'] == 'بدون معرّف وحيد'


def test_ligne_sans_identifiant_acceptee_sur_demande():
    plan = imp.analyser(classeur([
        ENTETE,
        ['', 'عريف', 'زياد', 'البوهلالي', '', ''],
    ]), [], accepter_sans_identifiant=True)
    assert len(plan['a_ajouter']) == 1
    assert plan['rejetes'] == []


def test_identifiant_repete_dans_le_fichier_compte_pour_un():
    plan = imp.analyser(classeur([
        ENTETE,
        ['77001', 'عريف', 'زياد', 'البوهلالي', '', ''],
        ['77001', 'رقيب', 'البوهلالي', 'زياد', 'ميناء بنزرت', ''],
    ]), [])
    assert len(plan['a_ajouter']) == 1
    assert 'مكرّر داخل الملفّ' in plan['rejetes'][0]['motif']
    assert '2' in plan['rejetes'][0]['motif']       # renvoie à la 1re occurrence


def test_doublon_de_nom_detecte_quand_lidentifiant_manque():
    plan = imp.analyser(classeur([
        ENTETE,
        ['', 'عريف', 'أنيس', 'زيتوني', '', ''],
        ['', 'رقيب', 'أنيس', 'زيتوني', '', ''],
    ]), [], accepter_sans_identifiant=True)
    assert len(plan['a_ajouter']) == 1
    assert len(plan['rejetes']) == 1


def test_deja_en_liste_par_identifiant_refuse():
    plan = imp.analyser(classeur([
        ENTETE,
        ['77001', 'وكيل', 'زياد', 'البوهلالي', 'ميناء بنزرت', ''],
    ]), [fiche_db(5, '77001')])
    assert plan['a_ajouter'] == []
    assert plan['rejetes'][0]['motif'] == 'مسجّل مسبقا في القائمة'


# ─── compléter, jamais écraser ───────────────────────────────────────────────

def test_completion_ne_remplit_que_les_خانات_vides():
    """Le cœur de l'affaire : ce qu'un agent a saisi prime sur le fichier."""
    existante = fiche_db(5, '77001', grade='عريف', lieu_travail='')
    plan = imp.analyser(classeur([
        ENTETE,
        ['77001', 'وكيل', 'زياد', 'البوهلالي', 'ميناء بنزرت',
         'الإدارة الجهويّة للدّيوانة ببنزرت'],
    ]), [existante], completer_existantes=True)

    assert plan['a_ajouter'] == []
    numero, mkow_id, partiel = plan['a_completer'][0]
    assert mkow_id == 5
    assert partiel['lieu_travail'] == 'ميناء بنزرت'
    assert partiel['jiha_marjiiya'] == 'الإدارة الجهويّة للدّيوانة ببنزرت'
    # الرتبة est déjà saisie : le « وكيل » du fichier ne passe pas.
    assert 'grade' not in partiel
    # ni le لقب ni l'اسم ne sont réécrits
    assert 'nom' not in partiel and 'prenom' not in partiel


def test_completion_sans_rien_a_completer_est_dite():
    existante = fiche_db(5, '77001', grade='عريف', lieu_travail='ميناء بنزرت')
    plan = imp.analyser(classeur([
        ENTETE,
        ['77001', 'عريف', 'زياد', 'البوهلالي', 'ميناء بنزرت', ''],
    ]), [existante], completer_existantes=True)
    assert plan['a_completer'] == []
    assert plan['rejetes'][0]['motif'] == 'مسجّل مسبقا ولا خانة فارغة تُستكمل'


# ─── النّموذج ────────────────────────────────────────────────────────────────

def test_modele_est_un_xlsx_lisible_de_droite_a_gauche():
    from openpyxl import load_workbook
    octets = imp.construire_modele().getvalue()
    assert octets[:2] == b'PK'                      # un .xlsx est un zip
    ws = load_workbook(BytesIO(octets)).active
    assert ws.title == 'المكوّنون'
    assert ws.sheet_view.rightToLeft is True
    assert [c.value for c in ws[1]] == list(imp.COLONNES_MODELE)


def test_le_modele_se_relit_lui_meme():
    """L'épreuve qui compte : ce que la منظومة propose, elle sait le relire."""
    octets = imp.construire_modele(
        [['77001', 'عريف', 'زياد', 'البوهلالي', 'ميناء بنزرت',
          'الإدارة الجهويّة للدّيوانة ببنزرت', '20 000 000']]).getvalue()
    plan = imp.analyser(BytesIO(octets), [])
    assert plan['ligne_entete'] == 1
    assert plan['colonnes_ignorees'] == []
    assert len(plan['a_ajouter']) == 1
    _, fiche = plan['a_ajouter'][0]
    assert fiche['identifiant_unique'] == '77001'
    assert fiche['nom'] == 'زياد'
    assert fiche['prenom'] == 'البوهلالي'
    assert fiche['telephone_gsm'] == '20 000 000'


# ─── الصّفحة ─────────────────────────────────────────────────────────────────

def _poster(client, lignes, **options):
    donnees = {'_csrf': 'jeton-de-test',
               'fichier': (classeur(lignes), 'mkowin.xlsx')}
    for cle, actif in options.items():
        if actif:
            donnees[cle] = '1'
    return client.post('/mkowin/importer', data=donnees,
                       content_type='multipart/form-data')


def test_page_importation_saffiche(db, client):
    rep = client.get('/mkowin/importer')
    assert rep.status_code == 200
    corps = rep.data.decode('utf-8')
    assert 'توريد' in corps
    assert 'المعرّف الوحيد' in corps
    assert '{{' not in corps                        # aucune variable non rendue


def test_modele_se_telecharge(db, client):
    rep = client.get('/mkowin/importer/namouthaj')
    assert rep.status_code == 200
    assert rep.data[:2] == b'PK'
    assert 'namouthaj_mkowin.xlsx' in rep.headers.get('Content-Disposition', '')


def test_sans_fichier_la_page_le_dit_sans_casser(db, client):
    rep = client.post('/mkowin/importer', data={'_csrf': 'jeton-de-test'},
                      content_type='multipart/form-data')
    assert rep.status_code == 200


def test_essai_n_ecrit_rien(db, client):
    """La تجربة montre et ne touche à rien : c'est sa raison d'être."""
    rep = _poster(client, [
        ENTETE,
        ['77001', 'عريف', 'زياد', 'البوهلالي', 'ميناء بنزرت', ''],
    ], essai=True)
    assert rep.status_code == 200
    assert 'البوهلالي' in rep.data.decode('utf-8')
    assert db.get_mkowin() == []


def test_versement_ecrit_vraiment(db, client):
    rep = _poster(client, [
        ENTETE,
        ['77001', 'عريف', 'زياد', 'البوهلالي', 'ميناء بنزرت',
         'الإدارة الجهويّة للدّيوانة ببنزرت'],
        ['77002', 'رقيب', 'سمير', 'قاسمي', 'مطار المنستير', ''],
    ], essai=False)
    assert rep.status_code == 200
    liste = db.get_mkowin()
    assert len(liste) == 2
    par_id = {m['identifiant_unique']: m for m in liste}
    assert par_id['77001']['nom'] == 'زياد'
    assert par_id['77001']['prenom'] == 'البوهلالي'
    assert par_id['77001']['jiha_marjiiya'] == 'الإدارة الجهويّة للدّيوانة ببنزرت'


def test_deux_versements_du_meme_fichier_ne_doublent_personne(db, client):
    lignes = [ENTETE, ['77001', 'عريف', 'زياد', 'البوهلالي', 'ميناء بنزرت', '']]
    _poster(client, lignes, essai=False)
    _poster(client, lignes, essai=False)
    assert len(db.get_mkowin()) == 1


def test_versement_ne_reecrit_pas_une_بطاقة_existante(db, client):
    """L'épreuve de la règle : ce qui est saisi ne se perd pas."""
    db.add_mkow({'identifiant_unique': '77001', 'nom': 'زياد',
                 'prenom': 'البوهلالي', 'grade': 'عريف', 'lieu_travail': ''})
    _poster(client, [
        ENTETE,
        ['77001', 'وكيل', 'زياد', 'البوهلالي', 'ميناء بنزرت', ''],
    ], essai=False, completer='1')
    fiche = db.get_mkowin()[0]
    assert fiche['grade'] == 'عريف'                 # الرتبة saisie : intacte
    assert fiche['lieu_travail'] == 'ميناء بنزرت'   # خانة vide : complétée


def test_fichier_illisible_ne_casse_pas_la_page(db, client):
    rep = client.post('/mkowin/importer',
                      data={'_csrf': 'jeton-de-test',
                            'fichier': (BytesIO(b'ceci n est pas un classeur'),
                                        'faux.xlsx')},
                      content_type='multipart/form-data')
    assert rep.status_code == 200
    assert db.get_mkowin() == []


def test_le_journal_garde_trace_du_versement(db, client):
    _poster(client, [
        ENTETE,
        ['77001', 'عريف', 'زياد', 'البوهلالي', '', ''],
    ], essai=False)
    conn = db.get_connection()
    try:
        lignes = conn.execute(
            "SELECT action, details FROM journal "
            "WHERE action LIKE '%توريد%'").fetchall()
    finally:
        conn.close()
    assert len(lignes) == 1
