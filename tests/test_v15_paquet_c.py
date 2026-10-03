# -*- coding: utf-8 -*-
"""v1.5 — الحزمة ج : exports Excel (D3), sauvegarde (S2), copie de programme
(P1), journal filtré (S1)."""
import io
import os
import sqlite3
import zipfile

from openpyxl import load_workbook

from tests.conftest import _client_flask
from core import importation, exports

JETON = {'X-CSRF-Token': 'jeton-de-test'}
MIME = exports.MIME_XLSX


def _classeur(r):
    assert r.status_code == 200 and r.mimetype == MIME, (r.status_code, r.mimetype)
    return load_workbook(io.BytesIO(r.data)).active


def _journal(db):
    c = db.get_connection()
    try:
        return [dict(r) for r in c.execute('SELECT * FROM journal ORDER BY id').fetchall()]
    finally:
        c.close()


def _mkow(db, **kw):
    base = {'grade': 'نقيب', 'nom': 'منى', 'prenom': 'بن عمر', 'identifiant_unique': '77001',
            'lieu_travail': 'مكتب الديوانة بفريانة', 'administration': 'الإدارة الجهوية',
            'telephone_gsm': '+216 20 000 000', 'cin': '01234567', 'banque': 'بنك',
            'num_compte': '12345678901234567890', 'notes': '=HYPERLINK("http://x")'}
    base.update(kw)
    db.add_mkow(base)


# ─── D3 : الأسماء ───────────────────────────────────────────────────────────

def test_export_mkowin(client, db):
    _mkow(db)
    _mkow(db, nom='سامي', prenom='الطرابلسي', identifiant_unique='77002', grade='عريف')
    r = client.get('/mkowin/exporter')
    assert 'attachment' in r.headers['Content-Disposition']
    ws = _classeur(r)
    assert ws.sheet_view.rightToLeft is True
    entetes = [c.value for c in ws[1]]
    assert entetes[:4] == ['المعرّف الوحيد', 'الرتبة', 'الاسم', 'اللقب']
    lignes = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert len(lignes) == 2
    i_notes = entetes.index('ملاحظات')
    note = [l for l in lignes if l[0] == '77001'][0][i_notes]
    assert note == '=HYPERLINK("http://x")'
    assert all(c.data_type != 'f' for row in ws.iter_rows() for c in row)   # aucune formule
    assert _journal(db)[-1]['action'] == 'تصدير قائمة الأسماء'


def test_export_mkowin_reimportable(client, db):
    """Aller-retour : les 12 premières colonnes sont reconnues par l'import."""
    _mkow(db)
    ws = _classeur(client.get('/mkowin/exporter'))
    entetes = [c.value for c in ws[1]]
    reconnus = [importation.champ_de_libelle(e) for e in entetes]
    assert reconnus[:12] == [ch for ch, _ in exports.COLONNES_MKOWIN[:12]]
    # les colonnes informatives ne sont prises pour AUCUN champ d'import
    assert all(r is None for r in reconnus[12:]), list(zip(entetes[12:], reconnus[12:]))
    flux = io.BytesIO(client.get('/mkowin/exporter').data)
    res = importation.analyser(flux, db.get_mkowin())
    txt = repr(res)
    assert '77001' in txt


def test_export_mkowin_vide(client):
    ws = _classeur(client.get('/mkowin/exporter'))
    assert ws.max_row == 1


# ─── D3 : سجلّ, دورات ───────────────────────────────────────────────────────

def test_export_registre(client, db):
    from tests.test_audit_v14e import _ouvrir_annee
    _ouvrir_annee(db, 2026)
    lid = db.save_programme('interne', 'أكتوبر', 2026, [{
        'titre': 'دورة', 'grade': 'مقدم', 'nom_formateur': 'س ع', 'lieu_travail': 'x',
        'date_formation': '2026-10-21', 'periode': '', 'lieu_formation': 'قاعة'}], 'n', 't')
    db.verrouiller_lettre(lid, 'interne')
    ws = _classeur(client.get('/registre/interne/exporter?annee=2026'))
    assert 'سنة 2026' in ws['A1'].value
    assert [c.value for c in ws[2]][:2] == ['الرقم', 'المرجع']
    assert ws['A3'].value == '0001'
    # type inconnu → interne ; année inconnue → année ouverte
    assert client.get('/registre/xxx/exporter?annee=1900').status_code == 200


def test_export_dorrat(client, db):
    from tests.test_audit_v14e import _dorra_finalisee
    _dorra_finalisee(db)
    ws = _classeur(client.get('/programmes/exporter'))
    ligne = [c.value for c in ws[2]]
    assert ligne[0] == 'تحرير المحاضر' and 'مسجّلة نهائيًّا — مقفلة' in ligne and 'المستحقّات غير منجزة' in ligne


# ─── S2 : sauvegarde ────────────────────────────────────────────────────────

def test_sauvegarde_telechargee_coherente(client, db):
    _mkow(db)
    r = client.get('/parametres/sauvegarde/telecharger')
    assert r.status_code == 200 and r.mimetype == 'application/zip'
    assert 'sauvegarde_formation_' in r.headers['Content-Disposition']
    z = zipfile.ZipFile(io.BytesIO(r.data))
    assert set(z.namelist()) == {'formation.db', 'LISEZ-MOI.txt'}
    chemin = os.path.join(os.path.dirname(db.DB_PATH), 'verif.db')
    open(chemin, 'wb').write(z.read('formation.db'))
    c = sqlite3.connect(chemin)
    assert c.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
    assert c.execute("SELECT nom FROM mkowin").fetchone()[0] == 'منى'
    c.close()
    assert _journal(db)[-1]['action'] == 'تنزيل نسخة احتياطية'


def test_sauvegarde_reservee_admin(db, monkeypatch):
    db.update_config('installation_faite', '1')
    c = _client_flask(db, monkeypatch, 'user')
    assert c.get('/parametres/sauvegarde/telecharger').status_code == 403
    assert c.post('/parametres/sauvegarde/creer', data={'_csrf': 'jeton-de-test'}).status_code == 403


def test_sauvegarde_creee_et_affichee(client, db):
    r = client.post('/parametres/sauvegarde/creer', data={'_csrf': 'jeton-de-test'})
    assert r.status_code == 302
    dossier = os.path.join(os.path.dirname(db.DB_PATH), 'backups')
    assert len(os.listdir(dossier)) == 1
    html = client.get('/parametres').get_data(as_text=True)
    assert 'عدد النسخ المحفوظة: <strong>1</strong>' in html
    assert 'telecharger' in html


# ─── P1 : copie de programme ────────────────────────────────────────────────

def _programme_source(db):
    from tests.test_audit_v14e import _dorra_finalisee
    return _dorra_finalisee(db)


def test_copie_programme(client, db):
    lid, fid = _programme_source(db)
    r = client.post(f'/lettre/{lid}/dupliquer', data={'mois': 'نوفمبر', 'annee': '2026',
                                                     '_csrf': 'jeton-de-test'})
    assert r.status_code == 302 and 'modifier=' in r.headers['Location']
    nouveau = int(r.headers['Location'].split('modifier=')[1])
    d = db.get_lettre_detail(nouveau)
    src = db.get_lettre_detail(lid)
    assert d['lettre']['mois'] == 'نوفمبر' and d['lettre']['annee'] == 2026
    assert not d['lettre']['verrouille'] and not d['lettre']['ref_complet']
    assert d['lettre']['type'] == src['lettre']['type']
    f, g = d['formations'][0], src['formations'][0]
    for k in ('titre', 'grade', 'nom_formateur', 'lieu_travail', 'lieu_formation', 'periode'):
        assert f[k] == g[k], k
    assert f['date_formation'] == ''
    c = db.get_connection()
    assert c.execute('SELECT COUNT(*) FROM participants WHERE lettre_id=?', (nouveau,)).fetchone()[0] == 0
    c.close()
    # la source n'a pas bougé
    assert db.get_lettre_detail(lid)['lettre']['verrouille']
    assert any(p['id'] == nouveau for p in db.get_programmes_inacheves())
    assert _journal(db)[-1]['action'] == 'نسخ برنامج تكوين'
    # l'écran d'édition s'ouvre bien sur la copie
    assert client.get(f'/lettre/nouvelle?modifier={nouveau}').status_code == 200


def test_copie_programme_refus(client, db):
    lid, _ = _programme_source(db)
    for data in ({'mois': 'شهر', 'annee': '2026'}, {'mois': 'نوفمبر', 'annee': 'x'},
                 {'mois': 'نوفمبر', 'annee': '1800'}, {}):
        data['_csrf'] = 'jeton-de-test'
        n = len(db.get_programmes_inacheves())
        r = client.post(f'/lettre/{lid}/dupliquer', data=data)
        assert r.status_code == 302 and len(db.get_programmes_inacheves()) == n
    r = client.post('/lettre/99999/dupliquer', data={'mois': 'ماي', 'annee': '2027', '_csrf': 'jeton-de-test'})
    assert r.status_code == 302 and 'programmes' in r.headers['Location']


def test_copie_refuse_une_mraslat_hurra(db):
    lid = db.save_programme('interne', 'ماي', 2026, [], 'n', 't')
    c = db.get_connection(); c.execute("UPDATE lettres SET categorie='libre' WHERE id=?", (lid,)); c.commit(); c.close()
    assert db.dupliquer_programme(lid, 'ماي', 2027) is None


def test_copie_formulaire_affiche(client, db):
    _programme_source(db)
    html = client.get('/programmes').get_data(as_text=True)
    assert 'dupliquer' in html and 'نسخ هذا البرنامج لشهر آخر' in html and 'data-confirm=' in html


# ─── S1 : journal filtré ────────────────────────────────────────────────────

def _remplir_journal(db):
    c = db.get_connection()
    for h, u, a in (('2026-01-05 10:00:00', 'admin', 'ترقيم مراسلة'),
                    ('2026-02-10 11:00:00', 'agent', 'حذف جهة مرجعيّة'),
                    ('2026-03-15 12:00:00', 'admin', 'تصدير قائمة الأسماء')):
        c.execute('INSERT INTO journal (horodatage, utilisateur, action, cible, details) VALUES (?,?,?,?,?)',
                  (h, u, a, '', ''))
    c.commit(); c.close()


def test_journal_filtres(db):
    _remplir_journal(db)
    f = db.get_journal_filtre
    assert [e['action'] for e in f(utilisateur='agent')] == ['حذف جهة مرجعيّة']
    assert len(f(du='2026-02-01', au='2026-03-31')) == 2
    assert len(f(au='2026-01-05')) == 1                      # borne incluse
    assert [e['action'] for e in f(q='تصدير')] == ['تصدير قائمة الأسماء']
    assert f(q="%' OR 1=1 --") == []                         # paramètres liés
    assert set(db.utilisateurs_du_journal()) >= {'admin', 'agent'}


def test_journal_page_et_export(client, db):
    _remplir_journal(db)
    html = client.get('/journal?utilisateur=agent').get_data(as_text=True)
    assert 'حذف جهة مرجعيّة' in html and 'ترقيم مراسلة' not in html
    assert 'عمليّة مطابقة للبحث' in html
    html = client.get('/journal?du=2030-01-01').get_data(as_text=True)
    assert 'لا توجد عمليّات مطابقة' in html
    html = client.get('/journal?du=nimportequoi').get_data(as_text=True)     # date invalide ignorée
    assert 'ترقيم مراسلة' in html
    ws = _classeur(client.get('/journal/exporter?utilisateur=admin'))
    actions = [r[2].value for r in ws.iter_rows(min_row=2)]
    assert 'ترقيم مراسلة' in actions and 'حذف جهة مرجعيّة' not in actions


def test_journal_export_reserve_admin(db, monkeypatch):
    db.update_config('installation_faite', '1')
    c = _client_flask(db, monkeypatch, 'user')
    assert c.get('/journal/exporter').status_code == 403


def test_formulaire_conserve_une_valeur_hors_liste():
    """Un titre ou un lieu absent des listes (مادة supprimée, programme copié)
    reste sélectionné au lieu d'être silencieusement effacé."""
    from tests.conftest import BASE_DIR
    from tests.conftest import js_programme
    s = open(os.path.join(BASE_DIR, 'templates', 'nouvelle_lettre.html'), encoding='utf-8').read() + js_programme()
    assert "_optionHorsListe(selected, MAWAD_INIT.map(m => m.titre))" in s
    assert '_optionHorsListe(selected, lieux)' in s
    assert 'غير موجود في القائمة' in s
