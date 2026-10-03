# -*- coding: utf-8 -*-
"""V3 — تصنيف الدّورات et module التّقارير.

Phase 1 : les 4 champs de تصنيف (نمط, مستوى, تعاون, خارج المخطّط) sont
enregistrés, relus et conservés par l'autosauvegarde ; ils n'ont AUCUN effet
sur le déroulement d'un برنامج."""

import os

from core import classification as cls

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

DORRA = {'titre': 'تحرير المحاضر', 'grade': 'مقدم', 'nom_formateur': 'البوهلالي زياد',
         'lieu_travail': 'الإدارة الجهوية للديوانة بالقصرين',
         'date_formation': '2026-10-21', 'periode': '',
         'lieu_formation': 'مركز التكوين الجهوي بالقصرين'}


def _classif(db, lettre_id):
    return [{k: f[k] for k in cls.CHAMPS} for f in db.get_lettre_detail(lettre_id)['formations']]


# ═══ Schéma ══════════════════════════════════════════════════════════════════

def test_schema_20_colonnes_et_table(db):
    conn = db.get_connection()
    cols = {c['name']: c for c in conn.execute('PRAGMA table_info(formations)')}
    for c in cls.CHAMPS:
        assert c in cols
    tables = {r['name'] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert 'parametres_annuels' in tables
    conn.close()


def test_anciennes_dorrat_prennent_les_valeurs_par_defaut(db, programme):
    """Une ligne insérée sans les nouvelles colonnes (cas des دورات V2)."""
    conn = db.get_connection()
    conn.execute("INSERT INTO formations (lettre_id, ordre, titre) VALUES (?, 9, 'قديمة')",
                 (programme['lettre_id'],))
    conn.commit()
    r = conn.execute("SELECT * FROM formations WHERE titre='قديمة'").fetchone()
    conn.close()
    assert (r['mode_formation'], r['niveau_formation'], r['cooperation'], r['hors_plan']) == \
        ('حضوري', 'جهوي', '', 0)


def test_migration_idempotente(db):
    from core.db import schema
    db.init_db()
    db.init_db()
    assert db.version_schema() == schema.SCHEMA_VERSION >= 20


# ═══ Normalisation ═══════════════════════════════════════════════════════════

def test_normaliser_defauts_et_valeurs_invalides():
    assert cls.normaliser({}) == {'mode_formation': 'حضوري', 'niveau_formation': 'جهوي',
                                  'cooperation': '', 'hors_plan': 0}
    n = cls.normaliser({'mode_formation': 'xxx', 'niveau_formation': 'مختص',
                        'cooperation': 'دولي', 'hors_plan': 'on'})
    assert n == {'mode_formation': 'حضوري', 'niveau_formation': 'مختص',
                 'cooperation': 'دولي', 'hors_plan': 1}
    assert cls.normaliser({'cooperation': True, 'hors_plan': False})['cooperation'] == 'وطني'


def test_normaliser_cle_absente_garde_l_ancienne_valeur():
    anc = {'mode_formation': 'عن بعد', 'niveau_formation': 'مركزي',
           'cooperation': 'وطني', 'hors_plan': 1}
    assert cls.normaliser({}, anc) == anc
    # clé présente : elle l'emporte
    assert cls.normaliser({'hors_plan': 0}, anc)['hors_plan'] == 0


def test_categorie_rapport():
    assert cls.categorie_rapport({}) == cls.CATEGORIE_PLAN
    assert cls.categorie_rapport({'cooperation': 'وطني'}) == cls.CATEGORIE_COOPERATION
    assert cls.categorie_rapport({'cooperation': 'وطني', 'hors_plan': 1}) == cls.CATEGORIE_HORS_PLAN


# ═══ Enregistrement ══════════════════════════════════════════════════════════

def test_save_programme_enregistre_la_classification(db):
    lid = db.save_programme('interne', 'أكتوبر', 2026, [
        dict(DORRA, mode_formation='عن بعد', niveau_formation='مختص',
             cooperation='دولي', hors_plan=1),
        dict(DORRA, date_formation='2026-10-22'),
    ], 'حازم مسعودي', 'النقيب')
    assert _classif(db, lid) == [
        {'mode_formation': 'عن بعد', 'niveau_formation': 'مختص', 'cooperation': 'دولي', 'hors_plan': 1},
        {'mode_formation': 'حضوري', 'niveau_formation': 'جهوي', 'cooperation': '', 'hors_plan': 0},
    ]


def test_update_programme_modifie_puis_conserve(db, programme):
    lid = programme['lettre_id']
    assert db.update_programme(lid, 'interne', 'أكتوبر', 2026,
                               [dict(DORRA, niveau_formation='مركزي', hors_plan=1)])
    assert _classif(db, lid)[0]['niveau_formation'] == 'مركزي'
    # payload d'un ancien onglet (sans les clés) : la classification survit
    assert db.update_programme(lid, 'interne', 'أكتوبر', 2026, [dict(DORRA)])
    assert _classif(db, lid)[0] == {'mode_formation': 'حضوري', 'niveau_formation': 'مركزي',
                                    'cooperation': '', 'hors_plan': 1}


def test_save_lettre_enregistre_la_classification(db):
    lid = db.save_lettre('interne', '', 0, 'أكتوبر', 2026,
                         [dict(DORRA, mode_formation='عن بعد', cooperation='وطني')],
                         'حازم مسعودي', 'النقيب')
    c = _classif(db, lid)[0]
    assert c['mode_formation'] == 'عن بعد' and c['cooperation'] == 'وطني'


def test_route_enregistrer_et_autosave(client, db):
    r = client.post('/lettre/enregistrer', headers={'X-CSRF-Token': 'jeton-de-test'}, json={
        'type': 'interne', 'mois': 'أكتوبر', 'annee': 2026,
        'formations': [dict(DORRA, mode_formation='عن بعد', niveau_formation='مختص',
                            cooperation='وطني', hors_plan=0)]})
    assert r.status_code == 200
    lid = r.get_json()['lettre_id']
    assert _classif(db, lid)[0]['niveau_formation'] == 'مختص'
    # la page d'édition renvoie la classification au JavaScript
    html = client.get(f'/lettre/nouvelle?modifier={lid}').get_data(as_text=True)
    assert '"niveau_formation": "\\u0645\\u062e\\u062a\\u0635"' in html or 'مختص' in html


def test_classification_sans_effet_sur_les_etapes(db):
    """Les étapes d'une دورة « عن بعد / خارج المخطّط » sont identiques."""
    a = db.save_programme('interne', 'أكتوبر', 2026, [dict(DORRA)], 'x', 'y')
    b = db.save_programme('interne', 'أكتوبر', 2026,
                          [dict(DORRA, mode_formation='عن بعد', hors_plan=1)], 'x', 'y')
    fa = db.get_lettre_detail(a)['formations'][0]
    fb = db.get_lettre_detail(b)['formations'][0]
    for k in ('titre', 'date_formation', 'date_fin', 'periode', 'lieu_formation'):
        assert fa[k] == fb[k]


# ═══ Interface ═══════════════════════════════════════════════════════════════

def test_interface_colonne_classification():
    tpl = open(os.path.join(BASE, 'templates', 'nouvelle_lettre.html'), encoding='utf-8').read()
    js = open(os.path.join(BASE, 'static', 'js', 'programme', '01_etape1_et_2.js'),
              encoding='utf-8').read()
    assert '>التصنيف</th>' in tpl
    assert 'function celluleClassification' in js and '...classificationDe(tr)' in js
    for c in ('f-mode', 'f-niveau', 'f-coop', 'f-coop-type', 'f-hors-plan'):
        assert c in js
    # noms de radios propres à chaque ligne
    assert 'name="${nom}_${n}"' in js


# ═══ Phase 2 — إعدادات التّقارير ═══════════════════════════════════════════

import json as _json

from core import rapports_defauts as rd


def _client_user(db, monkeypatch):
    from tests.conftest import _client_flask
    db.update_config('installation_faite', '1')
    return _client_flask(db, monkeypatch, 'user')


def test_textes_par_defaut_suivent_le_centre(db):
    config = db.get_config()
    assert rd.texte_legal(config) == rd.TEXTE_LEGAL_DEFAUT
    assert rd.fasl(config) == rd.FASL_DEFAUT
    assert rd.amr_court(rd.TEXTE_LEGAL_DEFAUT) == 'الأمر عدد 929 لسنة 2022'
    t = rd.taches_defaut(config)
    assert '{' not in t and t.count('\n- ') == 5
    s = rd.structure_defaut(config)
    assert 'الأمر عدد 929 لسنة 2022' in s and rd.FASL_DEFAUT in s
    intro, puces = rd.decouper_puces(t)
    assert intro.startswith('يشرف') and len(puces) == 5


def test_normaliser_rh():
    lignes = [{'niveau': '', 'off_sup': '2', 'homme': 'x', 'femme': -3},
              {'niveau': '', 'officier': ''},          # entièrement vide → écartée
              'pas un dict',
              {'niveau': 'مركزي', 'sous_off': 5000}]
    n = rd.normaliser_rh(lignes)
    assert len(n) == 2
    assert n[0]['niveau'] == 'جهوي' and n[0]['off_sup'] == 2
    assert n[0]['homme'] == 0 and n[0]['femme'] == 0
    assert n[1]['sous_off'] == 999
    assert rd.total_rh(n) == 2 + 999
    assert rd.normaliser_rh([{'niveau': 'جهوي', 'officier': 1}] * 30).__len__() == rd.RH_MAX_LIGNES
    assert _json.loads(rd.serialiser_rh(lignes)) == n


def test_valeur_a_stocker():
    assert rd.valeur_a_stocker('  نصّ\r\nآخر ', 'نصّ\nآخر') == ''
    assert rd.valeur_a_stocker('معدّل', 'نصّ') == 'معدّل'
    assert rd.valeur_a_stocker(None, 'نصّ') == ''


def test_page_parametres_rapports_get(client):
    r = client.get('/parametres/rapports')
    assert r.status_code == 200
    html = r.get_data(as_text=True)
    for attendu in ('إعدادات التّقارير', 'rapport_texte_legal', 'rapport_taches',
                    'rapport_structure', 'id="rhTable"', 'id="annees"', 'ضابط.س'):
        assert attendu in html


def test_post_textes_et_rh(client, db):
    lignes = [{'niveau': 'جهوي', 'officier': 1, 'homme': 1, 'r_centre': 1, 'direct': 1}]
    r = client.post('/parametres/rapports', data={
        '_csrf': 'jeton-de-test', 'action': 'textes',
        'rapport_texte_legal': rd.TEXTE_LEGAL_DEFAUT,       # = défaut → ''
        'rapport_fasl': 'الفصل 16',
        'rapport_taches': rd.taches_defaut(db.get_config()),
        'rapport_structure': 'هيكل معدّل',
        'rapport_rh': _json.dumps(lignes), 'rapport_rh_note': ' ملاحظة '})
    assert r.status_code == 302
    c = db.get_config()
    assert c.get('rapport_texte_legal', '') == ''
    assert c['rapport_fasl'] == 'الفصل 16'
    assert c.get('rapport_taches', '') == ''            # suit toujours le défaut
    assert c['rapport_structure'] == 'هيكل معدّل'
    assert rd.lire_rh(c)[0]['officier'] == 1 and rd.total_rh(rd.lire_rh(c)) == 1
    assert c['rapport_rh_note'] == 'ملاحظة'
    assert 'الفصل 16' in rd.structure_defaut(c)


def test_post_rh_json_invalide_ne_touche_rien(client, db):
    db.update_config('rapport_rh', rd.serialiser_rh([{'niveau': 'جهوي', 'officier': 2}]))
    client.post('/parametres/rapports', data={'_csrf': 'jeton-de-test',
                                              'rapport_rh': '{pas du json'})
    assert rd.total_rh(rd.lire_rh(db.get_config())) == 2


def test_parametres_rapports_refuse_non_admin(db, monkeypatch):
    c = _client_user(db, monkeypatch)
    r = c.get('/parametres/rapports')
    assert r.status_code in (302, 403)
    r = c.post('/parametres/rapports', data={'_csrf': 'jeton-de-test',
                                             'rapport_fasl': 'x'})
    assert r.status_code in (302, 403)
    assert db.get_config().get('rapport_fasl', '') in ('', None)


def test_parametres_annuels_crud(client, db):
    assert db.get_nb_programmees(2026) is None or db.get_nb_programmees(2026) == 0
    r = client.post('/parametres/rapports', data={'_csrf': 'jeton-de-test', 'action': 'annee',
                                                  'annee': '2026', 'nb': '24'})
    assert r.status_code == 302 and r.headers['Location'].endswith('#annees')
    assert db.get_nb_programmees(2026) == 24
    client.post('/parametres/rapports', data={'_csrf': 'jeton-de-test', 'action': 'annee',
                                              'annee': '2026', 'nb': '30'})
    assert db.get_nb_programmees(2026) == 30
    for annee, nb in (('1999', '5'), ('2026', '-1'), ('abc', '3'), ('2026', '10000')):
        assert db.set_parametre_annuel(annee, nb) is False
    assert db.get_nb_programmees(2026) == 30
    client.post('/parametres/rapports', data={'_csrf': 'jeton-de-test',
                                              'action': 'supprimer_annee', 'annee': '2026'})
    assert not db.get_nb_programmees(2026)


def test_bandeau_accueil_rh(client, db):
    html = client.get('/accueil').get_data(as_text=True)
    assert 'تعمير الجدول' in html
    client.post('/parametres/rapports/rh-plus-tard', data={'_csrf': 'jeton-de-test'})
    assert 'تعمير الجدول' not in client.get('/accueil').get_data(as_text=True)
    db.update_config('rapport_rh_ignore', '')
    db.update_config('rapport_rh', rd.serialiser_rh([{'niveau': 'جهوي', 'officier': 1}]))
    assert 'تعمير الجدول' not in client.get('/accueil').get_data(as_text=True)


def test_bandeau_absent_pour_non_admin(db, monkeypatch):
    c = _client_user(db, monkeypatch)
    assert 'تعمير الجدول' not in c.get('/accueil').get_data(as_text=True)


# ═══ Phase 3 — moteur de calcul (core/rapports_moteur.py) ═══════════════════
# Données FICTIVES uniquement.

import pytest

from core import rapports_moteur as rm


def _dorra(titre, formateur, grade, date, **extra):
    d = dict(DORRA, titre=titre, nom_formateur=formateur, grade=grade, date_formation=date)
    d.update(extra)
    return d


def _jeu(db, mois='مارس', annee=2025, dorrat=None):
    """Crée un برنامج et renvoie la liste des ids de ses دورات."""
    lid = db.save_programme('interne', mois, annee, dorrat, 'مكوّن وهمي', 'النقيب')
    return lid, [f['id'] for f in db.get_lettre_detail(lid)['formations']]


def _sql(db, requete, args=()):
    conn = db.get_connection()
    conn.execute(requete, args)
    conn.commit()
    conn.close()


def _finaliser(db, fid):
    _sql(db, "INSERT OR IGNORE INTO memo_formations (formation_id) VALUES (?)", (fid,))
    _sql(db, "UPDATE memo_formations SET finalise_at='2025-03-30 10:00:00' WHERE formation_id=?",
         (fid,))


PARTS = [  # nom, grade, sexe, fiaa
    ('فلان الأوّل', 'مقدم', 'ذكر', 'أقلّ من 40 سنة'),
    ('فلان الثّاني', 'نقيب', 'أنثى', '40 سنة وأكثر'),
    ('فلان الثّالث', 'عريف', 'ذكر', ''),
    ('فلان الرّابع', 'رقيب', '', ''),
    ('فلانة الخامسة', 'السيدة', 'أنثى', ''),
]


@pytest.fixture
def jeu(db):
    lid, ids = _jeu(db, dorrat=[
        _dorra('دورة أ', 'مكوّن أ', 'مقدم', '2025-03-03'),
        _dorra('دورة ب', 'مكوّن ب', 'السيدة', '2025-03-10'),
        _dorra('دورة ج', 'مكوّن أ', 'مقدم', '2025-03-17'),
    ])
    a, b, c = ids
    # دورة أ : 3 jours (lun→mer), مخطّط ; ب : تعاون دولي, عن بعد ; ج : خارج المخطّط, non finalisée
    _sql(db, "UPDATE formations SET date_fin='2025-03-05' WHERE id=?", (a,))
    _sql(db, "UPDATE formations SET cooperation='دولي', mode_formation='عن بعد' WHERE id=?", (b,))
    _sql(db, "UPDATE formations SET hors_plan=1, cooperation='وطني' WHERE id=?", (c,))
    _finaliser(db, a)
    _finaliser(db, b)
    db.save_participants(lid, a, [{'nom_prenom': n, 'grade': g, 'sexe': s, 'fiaa_omria': f,
                                   'identifiant_unique': f'X{i}'}
                                  for i, (n, g, s, f) in enumerate(PARTS)])
    db.save_participants(lid, b, [{'nom_prenom': 'فلان الأوّل', 'grade': 'مقدم',
                                   'identifiant_unique': 'X0', 'sexe': 'ذكر'}])
    db.save_participants(lid, c, [{'nom_prenom': 'فلان آخر', 'grade': 'ملازم'}])
    # غياب : فلان الثّالث absent de la دورة أ
    conn = db.get_connection()
    for p in conn.execute('SELECT id, nom_prenom FROM participants WHERE formation_id=?', (a,)):
        conn.execute('INSERT INTO mustahaqqat_hodour (formation_id, participant_id, present) '
                     'VALUES (?,?,?)', (a, p['id'], 0 if p['nom_prenom'] == 'فلان الثّالث' else 1))
    conn.execute("INSERT INTO mkowin (nom, prenom, grade, sexe) VALUES ('مكوّن', 'أ', 'مقدم', 'ذكر')")
    conn.commit()
    conn.close()
    return {'lettre_id': lid, 'ids': ids}


def test_bornes_periode_libelles():
    assert rm.bornes_periode(2025, 'trimestre', 1)['libelle'] == 'الثّلاثي الأوّل من سنة 2025'
    s2 = rm.bornes_periode(2025, 'semestre', 2)
    assert s2['libelle'] == 'السّداسي الثّاني من سنة 2025'
    assert (s2['debut'], s2['fin'], len(s2['mois'])) == ('2025-07-01', '2025-12-31', 6)
    an = rm.bornes_periode(2025, 'annee')
    assert an['libelle'] == 'سنة 2025' and len(an['mois']) == 12
    pl = rm.bornes_periode(type_periode='plage', debut='2025-02-01', fin='2025-02-28')
    assert pl['libelle'] == 'الفترة من 01/02/2025 إلى 28/02/2025' and pl['annee'] == 2025


@pytest.mark.parametrize('kw', [
    {'annee': 1999, 'type_periode': 'annee'},
    {'annee': 2025, 'type_periode': 'trimestre', 'n': 5},
    {'annee': 2025, 'type_periode': 'semestre', 'n': 0},
    {'type_periode': 'plage', 'debut': '2025-03-01', 'fin': '2025-02-01'},
    {'type_periode': 'plage', 'debut': 'xx', 'fin': '2025-02-01'},
    {'annee': 2025, 'type_periode': 'inconnu'},
])
def test_bornes_periode_invalides(kw):
    with pytest.raises(ValueError):
        rm.bornes_periode(**kw)


def test_meme_periode():
    p = rm.bornes_periode(2025, 'trimestre', 2)
    q = rm.meme_periode(p, 2)
    assert q['annee'] == 2023 and q['mois'] == p['mois']
    pl = rm.meme_periode(rm.bornes_periode(type_periode='plage', debut='2024-02-29',
                                           fin='2024-03-10'), 1)
    assert (pl['debut'], pl['fin']) == ('2023-02-28', '2023-03-10')


def test_evolution():
    assert rm.evolution(12, 10) == 20.0
    assert rm.evolution(5, 10) == -50.0
    assert rm.evolution(3, 0) is None


def test_categories_grades(db):
    conn = db.get_connection()
    t = rm._table_groupes(conn)
    conn.close()
    cp = rm.categorie_participant
    assert [cp(g, t) for g in ('العقيد', 'نقيب', 'الملازم أول', 'عريف', 'رقيب أول', 'السيد')] == \
        [rm.P_OFFICIERS, rm.P_OFFICIERS, rm.P_OFFICIERS, rm.P_SOUS_OFF, rm.P_RAQIB, rm.P_AUTRES]
    cf = rm.categorie_formateur
    assert [cf(g, t) for g in ('مقدم', 'الرائد', 'ملازم', 'الوكيل', 'رقيب', 'السيدة')] == \
        [rm.F_SUPERIEURS, rm.F_SUPERIEURS, rm.F_SUBALTERNES, rm.F_SOUS_OFF, rm.F_SOUS_OFF,
         rm.F_CIVILS]


def test_calculer_periode_vide(db):
    r = rm.calculer(rm.bornes_periode(2025, 'annee'))
    assert r['activites']['total'] == 0 and r['participants']['total'] == 0
    assert r['presence']['global']['taux_absence'] == 0.0
    assert r['realisation']['taux'] is None
    assert len(r['par_mois']) == 12


def test_calculer_finalisees(jeu):
    r = rm.calculer(rm.bornes_periode(2025, 'trimestre', 1))
    assert r['total_dorrat_periode'] == 3 and r['non_finalisees'] == 1
    act = r['activites']
    assert act['total'] == 2 and act['jours'] == 4          # 3 + 1
    cats = {c['categorie']: c['activites'] for c in act['categories']}
    assert cats[cls.CATEGORIE_PLAN] == 1 and cats[cls.CATEGORIE_COOPERATION] == 1
    assert cats[cls.CATEGORIE_HORS_PLAN] == 0
    p = r['participants']
    assert p['total'] == 6 and p['uniques'] == 5
    assert p['par_categorie'] == {rm.P_OFFICIERS: 3, rm.P_SOUS_OFF: 1, rm.P_RAQIB: 1,
                                  rm.P_AUTRES: 1}
    assert p['par_sexe'] == {'ذكر': 3, 'أنثى': 2, rm.NON_DEFINI: 1}
    assert p['moyenne'] == 3.0
    g = r['presence']['global']
    assert (g['total'], g['absents'], g['taux_absence']) == (6, 1, round(100 / 6, 2))
    assert r['presence']['par_categorie'][rm.P_SOUS_OFF]['absents'] == 1
    assert r['presence']['dorrat_sans_hodour'] == 1          # دورة ب
    f = r['formateurs']
    assert f['total'] == 2
    assert f['par_categorie'][rm.F_SUPERIEURS] == 1 and f['par_categorie'][rm.F_CIVILS] == 1
    assert f['par_sexe'] == {'ذكر': 1, 'أنثى': 1, rm.NON_DEFINI: 0}   # fiche + السيدة
    assert f['sans_fiche'] == 1
    mars = next(m for m in r['par_mois'] if m['mois'] == 'مارس')
    assert (mars['activites'], mars['participants'], mars['absents']) == (2, 6, 1)
    assert {x['valeur']: x['activites'] for x in r['par_mode']} == {'حضوري': 1, 'عن بعد': 1}
    assert {x['valeur']: x['activites'] for x in r['par_cooperation']} == {'وطني': 0, 'دولي': 1}


def test_calculer_toutes(jeu):
    r = rm.calculer(rm.bornes_periode(2025, 'trimestre', 1), critere='toutes')
    assert r['activites']['total'] == 3
    cats = {c['categorie']: c['activites'] for c in r['activites']['categories']}
    assert cats[cls.CATEGORIE_HORS_PLAN] == 1                # hors_plan prioritaire
    assert r['participants']['total'] == 7
    # مكوّن أ : 2 دورات
    assert r['par_formateur'][0]['activites'] == 2


def test_calculer_hors_periode(jeu):
    assert rm.calculer(rm.bornes_periode(2025, 'trimestre', 2))['activites']['total'] == 0
    pl = rm.bornes_periode(type_periode='plage', debut='2025-03-08', fin='2025-03-31')
    assert rm.calculer(pl)['activites']['total'] == 1         # دورة ب seulement


def test_taux_realisation(jeu, db):
    assert db.set_parametre_annuel(2025, 4)
    r = rm.calculer(rm.bornes_periode(2025, 'annee'))
    assert r['realisation'] == {'annee': 2025, 'programmees': 4, 'realisees_plan': 1,
                                'taux': 25.0}


def test_heures_seulement_si_acheve(jeu, db):
    a = jeu['ids'][0]
    _sql(db, "INSERT INTO mustahaqqat (formation_id, lettre_id, etat, heures, montant) "
             "VALUES (?,?,'hodour',6,90)", (a, jeu['lettre_id']))
    per = rm.bornes_periode(2025, 'annee')
    assert rm.calculer(per)['activites']['heures'] == 0
    _sql(db, "UPDATE mustahaqqat SET etat='acheve' WHERE formation_id=?", (a,))
    r = rm.calculer(per)
    assert r['activites']['heures'] == 6 and r['activites']['montant'] == 90
    fa = next(f for f in r['par_formateur'] if f['activites'] == 1 and f['heures'])
    assert fa['heures'] == 6


def test_comparer(jeu, db):
    _, ids = _jeu(db, annee=2024, dorrat=[_dorra('قديمة', 'مكوّن ج', 'نقيب', '2024-03-04')])
    _finaliser(db, ids[0])
    c = rm.comparer(rm.bornes_periode(2025, 'trimestre', 1), 3)
    assert len(c['precedents']) == 3
    assert c['precedents'][0]['annee'] == 2024 and c['precedents'][0]['activites'] == 1
    assert c['evolution_activites'] == 100.0
    assert c['evolution_absences'] is None                    # 0 absence en 2024
    assert rm.comparer(rm.bornes_periode(2025, 'annee'), 9)['precedents'].__len__() == 5


def test_moteur_hors_core_db():
    """La façade core/database.py n'est pas touchée par la phase 3."""
    assert not os.path.exists(os.path.join(BASE, 'core', 'db', 'rapports_moteur.py'))


# ═══ Phase 4 : التّقرير الكتابي ════════════════════════════════════════════════

import io
import json
import zipfile

from core import rapport_ecrit as rec

Q1 = {'type': 'trimestre', 'annee': 2025, 'n': 1}


def test_formats_nombres_et_tamyiz():
    assert (rec.n2(3), rec.n2(16), rec.n2(151), rec.n2(None)) == ('03', '16', '151', '00')
    assert (rec.dec(5.5), rec.dec(9.43), rec.dec(72.0), rec.dec(0)) == ('5.5', '9.43', '72', '0')
    assert rec.pct(58.27) == '58.27%'
    assert rec.nb(16, 'nashat') == '16 نشاطا'
    assert rec.nb(3, 'ghiyab') == '03 غيابات'
    assert rec.nb(0, 'mokawina') == '00 مكوّنات'
    assert rec.nb(11, 'ghiyab') == '11 غيابا'
    assert rec.nb(100, 'moucharik') == '100 مشارك'
    assert rec.nb(110, 'yawm') == '110 أيّام'
    assert rec.avec_lam('المركز الجهوي') == 'للمركز الجهوي'
    assert rec.avec_lam('مركز التّكوين') == 'لمركز التّكوين'


def test_contenu_textes_et_tableaux(jeu, db):
    p = rm.bornes_periode(2025, 'trimestre', 1)
    c = rec.contenu(p, db.get_config())
    t = c['textes']
    assert t['titre'].startswith('التّقرير الثّلاثي لنشاط')
    assert '02 نشاطا' in t['act_intro'] and 'التّعاون الوطني' in t['act_intro']
    assert '06 مشاركين' in t['part_intro']
    assert 'عدد 01 غيابا' in t['absences']
    assert '……%' in t['realisation']                     # pas de دورات مبرمجة
    assert any('المبرمجة' in a for a in c['alertes'])
    assert any('غير مختومة' in a for a in c['alertes'])
    # tableaux : en-têtes du modèle, entiers sur 2 chiffres
    cat = c['tableaux']['categories']['lignes']
    assert cat[1][0]['t'] == 'تصنيف النّشاط' and cat[-1][0]['t'] == 'المجموع العامّ'
    assert cat[-1][1]['t'] == '02'
    rh = c['tableaux']['rh']
    assert len(rh['largeurs']) == 13 and rh['lignes'][1][0]['vm'] == 'debut'
    assert [x['t'] for x in c['tableaux']['participants']['lignes'][1]][1:5] == \
        list(rm.CATEGORIES_PARTICIPANTS)
    # l'en-tête du rapport omet « الإدارة العامّة للدّيوانة » comme le modèle
    assert not any('الإدارة العامّـة' in l for l in c['entete'])


def test_brouillon_ne_garde_que_les_retouches(jeu, db):
    p = rm.bornes_periode(2025, 'trimestre', 1)
    auto = rec.contenu(p, db.get_config(), brouillon={})['auto']
    d = rec.brouillon_depuis_formulaire(
        {'intro1': auto['intro1'] + '\r\n', 'titre': 'عنوان جديد', 'difficultes': 'نقص المكوّنين'},
        auto)
    assert d == {'titre': 'عنوان جديد', 'difficultes': 'نقص المكوّنين'}
    db.update_config(rec.cle_brouillon(p), json.dumps(d, ensure_ascii=False))
    c = rec.contenu(p, db.get_config())
    assert c['textes']['titre'] == 'عنوان جديد' and c['modifies'] == ['difficultes', 'titre']
    assert ('puce', 'نقص المكوّنين', '▪') in rec.plan(c)
    assert ('pointilles',) in rec.plan(c)                 # المقترحات vides


def test_page_rapport_ecrit(client, jeu):
    r = client.get('/rapports/ecrit?type=trimestre&annee=2025&n=1')
    assert r.status_code == 200
    html = r.get_data(as_text=True)
    for fragment in ('التّقرير الكتابي', 'name="act_intro"', 'name="difficultes"',
                     'توزيع الأنشطة التّكوينيّة', 'btnPdf', 'btnWord'):
        assert fragment in html, fragment
    assert client.get('/rapports/ecrit').status_code == 200            # période par défaut
    assert client.get('/rapports/ecrit?type=plage&debut=x&fin=y').status_code == 200
    assert '/rapports/ecrit' in client.get('/statistiques').get_data(as_text=True)


def test_enregistrer_et_reinitialiser_brouillon(client, db, jeu):
    data = dict(Q1, _csrf='jeton-de-test', titre='تقرير معدّل', propositions='مقترح أوّل\nمقترح ثان')
    r = client.post('/rapports/ecrit/brouillon', data=data, headers={'X-Requested-With': 'fetch'})
    assert r.status_code == 200 and r.get_json()['ok']
    cle = 'rapport_brouillon_trimestre_2025_1'
    d = json.loads(db.get_config()[cle])
    assert d == {'titre': 'تقرير معدّل', 'propositions': 'مقترح أوّل\nمقترح ثان'}
    assert 'تقرير معدّل' in client.get('/rapports/ecrit?type=trimestre&annee=2025&n=1') \
        .get_data(as_text=True)
    r = client.post('/rapports/ecrit/brouillon', data=dict(Q1, _csrf='jeton-de-test',
                                                           action='reinitialiser'))
    assert r.status_code == 302 and db.get_config()[cle] == ''


def test_brouillon_exige_csrf(client, db, jeu):
    r = client.post('/rapports/ecrit/brouillon', data=dict(Q1, titre='x'))
    assert r.status_code in (400, 403)
    assert not db.get_config().get('rapport_brouillon_trimestre_2025_1')


def test_pdf_et_word(client, jeu):
    r = client.get('/rapports/ecrit/pdf?type=trimestre&annee=2025&n=1')
    assert r.status_code == 200 and r.data[:4] == b'%PDF' and r.mimetype == 'application/pdf'
    r = client.get('/rapports/ecrit/word?type=trimestre&annee=2025&n=1')
    assert r.status_code == 200 and 'attachment' in r.headers['Content-Disposition']
    z = zipfile.ZipFile(io.BytesIO(r.data))
    doc = z.read('word/document.xml').decode('utf-8')
    assert 'التّقرير الثّلاثي لنشاط' in doc and 'PAGE' in z.read('word/footer1.xml').decode()
    assert '<w:ind w:firstLine=' in doc and 'w:fill="8DB3E2"' in doc


def test_pdf_word_periode_vide(client, db):
    """Aucune دورة : le rapport sort quand même (zéros + lignes pointillées)."""
    for fmt in ('pdf', 'word'):
        r = client.get(f'/rapports/ecrit/{fmt}?type=annee&annee=2031')
        assert r.status_code == 200


def test_rapport_ecrit_refuse_anonyme(db, monkeypatch):
    import app as application
    c = application.app.test_client()
    assert c.get('/rapports/ecrit').status_code in (302, 401, 403)
    assert c.get('/rapports/ecrit/pdf').status_code in (302, 401, 403)


# ═══ Phase 5 : التّقرير البياني والتّحليلي ═════════════════════════════════════

from core import rapports_graphiques as rg


def test_donnees_graphiques(jeu, db):
    d = rg.donnees(rm.bornes_periode(2025, 'trimestre', 1), 2)
    assert [k['cle'] for k in d['kpis']][:3] == ['activites', 'jours', 'participants']
    assert d['kpis'][0]['valeur'] == 2 and d['kpis'][0]['evolution'] is None   # 2024 vide
    assert len(d['serie']) == 3 and d['serie'][-1]['annee'] == 2025
    ids = [g['id'] for s in d['sections'] for g in s['graphes']]
    for gid in ('cmp_act', 'mois_act', 'cat_act', 'part_sexe', 'pres_cat', 'form_top'):
        assert gid in ids
    for s in d['sections']:
        for g in s['graphes']:
            assert len(g['libelles']) == len(g['valeurs']), g['id']
            assert g['forme'] in ('barres_v', 'barres_h', 'pile')
    assert d['analyses'] and 'نشاطا' in d['analyses'][0]
    json.dumps(d, ensure_ascii=False)                         # sérialisable (page / export)


def test_analyses_periode_vide(db):
    d = rg.donnees(rm.bornes_periode(2031, 'annee'), 1)
    assert d['analyses'] == ['لا توجد أنشطة محتسبة في هذه الفترة.']


def test_page_rapports(client, jeu):
    r = client.get('/rapports?type=trimestre&annee=2025&n=1&annees=3')
    assert r.status_code == 200
    html = r.get_data(as_text=True)
    for fragment in ('التّقارير البيانيّة', 'data-graphe="cmp_act"', 'id="rgDonnees"',
                     'rapports_graphes.js', 'أهمّ الاستنتاجات', 'selected>3 سنوات'):
        assert fragment in html, fragment
    assert client.get('/rapports').status_code == 200
    assert client.get('/rapports?annees=99&toutes=1').status_code == 200
    assert 'url_for' not in client.get('/statistiques').get_data(as_text=True)
    assert 'href="/rapports"' in client.get('/statistiques').get_data(as_text=True)


def test_rapports_pdf_et_html(client, jeu):
    r = client.get('/rapports/pdf?type=trimestre&annee=2025&n=1')
    assert r.status_code == 200 and r.data[:4] == b'%PDF'
    # V3.0.1 : la « نسخة HTML تفاعليّة » est retirée (route et bouton)
    assert client.get('/rapports/html?type=trimestre&annee=2025&n=1').status_code == 404
    assert 'نسخة HTML' not in client.get('/rapports').get_data(as_text=True)
    assert client.get('/rapports/pdf?type=annee&annee=2031').status_code == 200


# ═══ Phase 6a : البطاقة التّقييميّة للمشارك ════════════════════════════════════

from core import fiche_participant as fp


def test_recherche_participant(jeu):
    r = fp.rechercher('فلان الأول')                     # sans شدّة ni hamza : trouvé quand même
    assert r and r[0]['cle'] == 'id:X0' and r[0]['dorrat'] == 2
    assert fp.rechercher('X3')[0]['nom'] == 'فلان الرّابع'
    assert fp.rechercher('') == [] and fp.rechercher('لا يوجد') == []


def test_fiche_participant(jeu):
    f = fp.fiche('id:X0')
    assert f['identite']['nom'] == 'فلان الأوّل' and f['identite']['categorie'] == 'ضبّاط'
    t = f['totaux']
    assert t['dorrat'] == 2 and t['closes'] == 2 and t['absences'] == 0
    assert t['taux_presence'] == 100.0 and t['mawadhaba'] == 'ممتازة'
    assert t['jours'] == 4                                 # 3 jours + 1 jour
    f3 = fp.fiche('id:X2')                                 # فلان الثّالث : absent
    assert f3['totaux']['absences'] == 1 and f3['totaux']['mawadhaba'] == 'ضعيفة'
    assert f3['dorrat'][0]['present'] is False and f3['totaux']['jours'] == 0
    autre = fp.fiche('nom:' + fp.normaliser('فلان آخر'))  # دورة non finalisée
    assert autre['totaux']['en_cours'] == 1 and autre['totaux']['taux_presence'] is None
    assert autre['dorrat'][0]['present'] is None
    assert fp.fiche('id:inconnu') is None
    assert fp.fiche('id:X0', 2031)['totaux']['dorrat'] == 0


def test_pages_fiche_participant(client, jeu):
    assert client.get('/rapports/participant').status_code == 200
    html = client.get('/rapports/participant?q=فلان').get_data(as_text=True)
    assert 'فلان الأوّل' in html and 'p=id:X0' in html
    html = client.get('/rapports/participant?p=id:X0').get_data(as_text=True)
    for fragment in ('مؤشّر المواظبة', 'سجلّ الدّورات', 'دورة أ', 'fpGraphes'):
        assert fragment in html, fragment
    j = client.get('/rapports/participant/recherche?q=فلان').get_json()
    assert len(j['resultats']) >= 5
    r = client.get('/rapports/participant/pdf?p=id:X0')
    assert r.status_code == 200 and r.data[:4] == b'%PDF'
    assert client.get('/rapports/participant/pdf?p=id:inconnu').status_code == 404
    assert client.get('/rapports/participant?p=id:inconnu').status_code == 200


# ═══ Phase 6b : التّقارير التّفصيليّة + Excel ══════════════════════════════════

from core import rapports_details as rdet


def test_tableaux_details(jeu):
    d = rdet.tableaux(rm.bornes_periode(2025, 'trimestre', 1))
    ids = [t['id'] for t in d['tableaux']]
    assert ids == ['resume', 'formateurs', 'dorrat', 'mois', 'sexe', 'fiaa', 'categorie',
                   'niveau', 'mode', 'cooperation']
    for t in d['tableaux']:
        for l in t['lignes'] + ([t['total']] if t['total'] else []):
            assert len(l) == len(t['entetes']), t['id']
    dorrat = next(t for t in d['tableaux'] if t['id'] == 'dorrat')
    assert len(dorrat['lignes']) == 2 and dorrat['total'][8] == 6      # 5 + 1 participants
    toutes = rdet.tableaux(rm.bornes_periode(2025, 'trimestre', 1), 'toutes')
    assert len(next(t for t in toutes['tableaux'] if t['id'] == 'dorrat')['lignes']) == 3


def test_page_et_excel_details(client, jeu):
    r = client.get('/rapports/details?type=trimestre&annee=2025&n=1&onglet=sexe')
    assert r.status_code == 200
    html = r.get_data(as_text=True)
    assert 'حسب المكوّنين' in html and 'id="pn-sexe"' in html and 'دورة أ' in html
    assert client.get('/rapports/details?onglet=inconnu').status_code == 200
    r = client.get('/rapports/details/excel?type=trimestre&annee=2025&n=1')
    assert r.status_code == 200 and r.data[:2] == b'PK'
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.data))
    assert len(wb.sheetnames) == 10 and wb[wb.sheetnames[0]].sheet_view.rightToLeft
    assert client.get('/rapports/details/excel?type=annee&annee=2031').status_code == 200


# ═══ Phase 6c : مؤشّرات الشّاشة الرّئيسيّة ═════════════════════════════════════

def test_kpis_accueil(jeu):
    from datetime import date
    k = rg.kpis_accueil(date(2025, 3, 31))
    assert k['annee'] == 2025 and not k['vide']
    assert k['indicateurs'][0]['valeur'] == 2 and k['indicateurs'][0]['evolution'] is None
    g = k['sections'][0]['graphes'][0]
    assert g['forme'] == 'barres_g' and len(g['libelles']) == 12
    assert [s['nom'] for s in g['series']] == ['2025', '2024'] and g['series'][0]['valeurs'][2] == 2
    assert rg.kpis_accueil(date(2031, 5, 1))['vide']
    assert rg.kpis_accueil(date(2028, 2, 29))['annee'] == 2028       # 29 février → 28 en 2027


def test_accueil_affiche_kpis(client, jeu, monkeypatch):
    from datetime import date
    reel = rg.kpis_accueil
    monkeypatch.setattr(rg, 'kpis_accueil', lambda aujourd_hui=None: reel(date(2025, 3, 31)))
    html = client.get('/accueil').get_data(as_text=True)
    assert 'id="accKpis"' in html and 'نشاط سنة 2025' in html


def test_accueil_resiste_a_une_erreur_kpis(client, monkeypatch):
    def casse(*a, **k):
        raise RuntimeError('panne simulée')
    monkeypatch.setattr(rg, 'kpis_accueil', casse)
    r = client.get('/accueil')
    assert r.status_code == 200 and 'id="accKpis"' not in r.get_data(as_text=True)
