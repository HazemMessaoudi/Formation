# -*- coding: utf-8 -*-
"""v1.5 — الحزمة د : لوحة القيادة (D2), التقرير السنوي (D1), graphiques,
recherche unifiée (D4)."""
import io
import re
from datetime import date, timedelta

from openpyxl import load_workbook

from core import graphiques
from core.recherche import normaliser, rechercher
from tests.test_mustahaqqat_qima import _dorra, _p, _jusquau_classement


def _dorra_payee(db):
    """Une دورة dont les المستحقّات sont approuvées : 4 h, 100.000 د."""
    fid = _dorra(db, [_p('أ', 'المقدم', '1'), _p('ب', 'المقدم', '2')])
    _jusquau_classement(db, fid, 'أ1')
    ok, info = db.confirmer_mustahaqqat(fid)
    assert ok and info['montant'] == 100.0 and info['heures'] == 4
    return fid


# ─── graphiques ─────────────────────────────────────────────────────────────

def test_maximum_rond():
    # effectifs : pas entiers (1, 2, 5, 10, 25… par graduation), 4 graduations
    assert [graphiques.maximum_rond(v, True) for v in (0, 2, 4, 5, 7, 11, 23, 47, 99, 101, 740)] == \
        [4, 4, 4, 8, 8, 20, 40, 80, 100, 200, 800]
    # montants : le pas 2.5 est admis
    assert [graphiques.maximum_rond(v) for v in (100, 137.5, 1234.5)] == [100, 200, 2000]


def test_graduations_entieres_pour_des_effectifs():
    g = graphiques.colonnes(['أ', 'ب'], [2, 1])
    assert [t['texte'] for t in g['graduations']] == ['0', '1', '2', '3', '4']


def test_colonnes_geometrie():
    g = graphiques.colonnes(['أ', 'ب', 'ج'], [0, 4, 2], unite='دورة')
    cols = g['colonnes']
    assert cols[0]['chemin'] == ''                             # zéro : pas de colonne
    assert cols[1]['etiquette'] and cols[2]['etiquette']       # V3.0.1 : toute colonne non nulle est écrite
    assert not cols[0]['etiquette']                            # zéro : pas d'étiquette
    assert cols[0]['x_bande'] > cols[1]['x_bande'] > cols[2]['x_bande']   # droite → gauche
    assert cols[1]['infobulle'] == 'ب: 4 دورة'
    assert not g['vide']
    assert graphiques.colonnes(['أ'], [0])['vide']


def test_colonnes_largeur_plafonnee():
    g = graphiques.colonnes(['أ'], [3])
    m = re.findall(r'[ML]([\d.]+),', g['colonnes'][0]['chemin'])
    xs = [float(x) for x in m]
    assert max(xs) - min(xs) <= graphiques.COLONNE_MAX + 0.01


# ─── D2 : لوحة القيادة ──────────────────────────────────────────────────────

def test_tableau_de_bord_compte_ce_qui_demande_attention(db):
    j = date(2026, 3, 15)
    # programme en retard (brouillon dont la دورة est passée)
    db.save_programme('interne', 'مارس', 2026, [{
        'titre': 'دورة فائتة', 'grade': 'مقدم', 'nom_formateur': 'س ع', 'lieu_travail': 'x',
        'date_formation': '2026-03-02', 'periode': '', 'lieu_formation': 'قاعة'}], 'n', 't')
    # دورة à venir dans 10 jours
    db.save_programme('interne', 'مارس', 2026, [{
        'titre': 'دورة قادمة', 'grade': 'مقدم', 'nom_formateur': 'س ع', 'lieu_travail': 'x',
        'date_formation': '2026-03-25', 'periode': '', 'lieu_formation': 'قاعة'}], 'n', 't')
    tb = db.donnees_tableau_de_bord(aujourdhui=j)
    assert [p['titres'] for p in tb['en_retard']] == ['دورة فائتة']
    assert tb['brouillons'] == 2
    assert [(d['titre'], d['dans_jours']) for d in tb['a_venir']] == [('دورة قادمة', 10)]
    assert tb['par_mois'][2]['mois'] == 'مارس' and tb['par_mois'][2]['dorrat'] == 2
    assert tb['dorrat_annee'] == 2


def test_tableau_de_bord_attente_mustahaqqat(db):
    fid = _dorra(db, [_p('أ', 'المقدم', '1')])
    assert db.donnees_tableau_de_bord()['attente_mustahaqqat'] == 1
    _jusquau_classement(db, fid, 'أ1')
    db.confirmer_mustahaqqat(fid)
    assert db.donnees_tableau_de_bord()['attente_mustahaqqat'] == 0


def test_page_tableau_de_bord(client, db):
    _dorra(db, [_p('أ', 'المقدم', '1')])
    db.journaliser('admin', 'عمليّة للتجربة')
    html = client.get('/dashboard').get_data(as_text=True)
    for fragment in ('stat-val', 'يتطلّب انتباهك', 'الدورات حسب الشهر', 'الدورات القادمة',
                     'class="graph__svg"', 'data-infobulle="فيفري: 1 دورة"', 'المعطيات في جدول',
                     'آخر العمليّات'):
        assert fragment in html, fragment


def test_tableau_de_bord_agent_sans_journal(db, monkeypatch):
    from tests.conftest import _client_flask
    db.update_config('installation_faite', '1')
    c = _client_flask(db, monkeypatch, 'user')
    db.journaliser('admin', 'عمليّة للتجربة')
    html = c.get('/dashboard').get_data(as_text=True)
    assert 'يتطلّب انتباهك' in html and 'آخر العمليّات' not in html


# ─── D1 : التقرير السنوي ────────────────────────────────────────────────────

def test_rapport_annuel_chiffres(db):
    _dorra_payee(db)
    _dorra(db, [_p('ج', 'الرقيب', '3')])                     # enregistrée, non payée
    annee = db.annee_registre()
    r = db.rapport_annuel(annee)
    t = r['totaux']
    assert t['dorrat'] == 2 and t['finalisees'] == 2 and t['participations'] == 3
    assert t['participants_uniques'] == 3 and t['dorrat_payees'] == 1
    assert t['heures'] == 4 and t['montant'] == 100.0          # seule la دورة approuvée
    fev = r['par_mois'][1]
    assert fev['mois'] == 'فيفري' and fev['dorrat'] == 2 and fev['montant'] == 100.0
    assert sum(m['dorrat'] for m in r['par_mois']) == t['dorrat']
    assert r['par_formateur'][0]['nom'] == 'زياد البوهلالي' and r['par_formateur'][0]['heures'] == 4
    assert r['par_madda'][0] == {'titre': 'تحرير المحاضر', 'dorrat': 2, 'participants': 3}
    assert db.rapport_annuel(annee - 5)['totaux']['dorrat'] == 0


def test_rapport_coherent_avec_les_mustahaqqat(db):
    """Le total du rapport = la somme affichée par l'écran des المستحقّات."""
    _dorra_payee(db)
    s = db.get_stats_mustahaqqat()
    r = db.rapport_annuel(db.annee_registre())
    assert r['totaux']['montant'] == float(s['montant_total'])
    assert r['totaux']['heures'] == s['heures_payees']


def test_rapport_ignore_les_mraslat_hurra(db):
    lid = db.save_programme('interne', 'ماي', 2026, [{
        'titre': 'x', 'grade': '', 'nom_formateur': '', 'lieu_travail': '', 'date_formation': '',
        'periode': '', 'lieu_formation': ''}], 'n', 't')
    c = db.get_connection(); c.execute("UPDATE lettres SET categorie='libre' WHERE id=?", (lid,)); c.commit(); c.close()
    t = db.rapport_annuel(2026)['totaux']
    assert t['programmes'] == 0 and t['dorrat'] == 0 and t['libres'] == 1


def test_page_rapport_annuel(client, db):
    _dorra_payee(db)
    annee = db.annee_registre()
    html = client.get(f'/statistiques/rapport-annuel?annee={annee}').get_data(as_text=True)
    assert f'التقرير السنوي <span dir="ltr">{annee}</span>' in html
    assert '100.000' in html and 'زياد البوهلالي' in html
    # V3.0.1 : التقرير السنوي est rangé sous التّقارير (plus un titre du menu)
    stats = client.get('/statistiques').get_data(as_text=True)
    assert 'التقرير السنوي' not in stats and 'href="/rapports"' in stats
    assert 'href="/statistiques/rapport-annuel' in client.get('/rapports').get_data(as_text=True)
    # année inconnue → la plus récente disponible
    assert client.get('/statistiques/rapport-annuel?annee=1800').status_code == 200


def test_export_rapport_annuel(client, db):
    _dorra_payee(db)
    annee = db.annee_registre()
    r = client.get(f'/statistiques/rapport-annuel/exporter?annee={annee}')
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ['ملخّص', 'حسب الشهر', 'حسب المكوّن', 'حسب المادّة', 'حسب الجنس']
    mois = wb['حسب الشهر']
    total = [c.value for c in mois[mois.max_row]]
    assert total[0] == 'المجموع' and total[1] == 1 and total[5] == 100.0


# ─── D4 : recherche unifiée ─────────────────────────────────────────────────

def test_normaliser():
    assert normaliser('  أَحْمَدُ  إبراهيم ') == 'احمد ابراهيم'
    assert normaliser('مدرسة') == normaliser('مدرسه')
    assert normaliser('مصطفى') == normaliser('مصطفي')


def test_recherche_groupes(db):
    db.add_mkow({'grade': 'نقيب', 'nom': 'منى', 'prenom': 'بن عمر', 'identifiant_unique': '77001'})
    db.add_madda({'titre': 'تقنيات التفتيش'}) if hasattr(db, 'add_madda') else None
    _dorra(db, [_p('أ', 'المقدم', '1')])
    g = {x['groupe']: x for x in rechercher('منى')}
    assert 'الأسماء' in g and g['الأسماء']['resultats'][0]['titre'] == 'منى بن عمر'
    assert rechercher('77001')[0]['resultats'][0]['cible'][0] == 'fiche_mkow'
    assert 'الدورات' in {x['groupe'] for x in rechercher('المحاضر')}
    assert 'المراسلات' in {x['groupe'] for x in rechercher('END')}
    # ET logique, ordre libre, sans حركات
    assert rechercher('عمر مُنى')[0]['resultats'][0]['titre'] == 'منى بن عمر'
    assert rechercher('منى زززز') == []
    assert rechercher('م') == [] and rechercher('') == []


def test_api_recherche(client, db):
    db.add_mkow({'grade': 'نقيب', 'nom': 'منى', 'prenom': 'بن عمر'})
    mid = db.get_mkowin()[0]['id']
    d = client.get('/api/recherche?q=منى').get_json()
    assert d['groupes'][0]['resultats'][0]['url'] == f'/mkowin/{mid}'
    assert client.get('/api/recherche?q=%25%27%22<script>').get_json()['groupes'] == []


def test_champ_recherche_present_sauf_mode_restreint(client, db):
    assert 'id="rechGlobaleInput"' in client.get('/dashboard').get_data(as_text=True)
    with client.session_transaction() as s:
        s['doit_changer_mdp'] = True
    c = db.get_connection(); c.execute("UPDATE users SET doit_changer_mdp=1 WHERE username='admin'"); c.commit(); c.close()
    assert 'id="rechGlobaleInput"' not in client.get('/changer-mot-de-passe').get_data(as_text=True)


# ─── Finitions v1.5 ─────────────────────────────────────────────────────────

def test_version_1_5():
    """Depuis la v1.6 le numéro a avancé : on vérifie seulement la cohérence."""
    from core import identite
    assert identite.VERSION_LABEL == f'الإصدار {identite.VERSION_APP}'
    # 1.0 : numérotation repartie de 1.0 à la mise en service
    assert all(x.isdigit() for x in identite.VERSION_APP.split('.'))


def test_ecran_de_connexion_sans_centre_code_en_dur(db):
    import app as application
    db.update_config('nom_centre', 'مركز التكوين الجهوي بسوسة')
    application.app.config.update(TESTING=True)
    html = application.app.test_client().get('/login').get_data(as_text=True)
    assert 'مركز التكوين الجهوي بسوسة' in html
    assert 'الإدارة الجهوية للديوانة – القصرين' not in html
